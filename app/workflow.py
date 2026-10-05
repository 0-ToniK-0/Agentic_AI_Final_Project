"""The agents and the two LangGraph workflows of Final_Notebook.ipynb, as an importable module.

- Delivery workflow (`builder`): an unclear request email -> clarification email -> customer answers ->
  approved specification -> plan, FDD/TDD, tests -> validation (reflexion loop) -> final approval -> publish.
- Support workflow (`support_builder`): a support email -> ticket with small tasks under the customer's
  support Feature -> Change Request when needed -> SharePoint -> acknowledgement email.

The graphs pause with interrupt() at every human gate; run.py answers them on the command line and compiles
the graphs with a checkpointer. The code follows the notebook cell by cell (sections 1-10).

Import this module after run.py has loaded .env and moved into the data folder: the clients read their
settings when the module is imported, and every runtime file is written in the current folder.
"""
import contextlib
import io
import json
import operator
import os
from datetime import date, timedelta
from pathlib import Path
from typing import Annotated, TypedDict

from langchain_core.messages import HumanMessage, SystemMessage
from langgraph.graph import StateGraph, START, END
from langgraph.types import Command, interrupt
from langsmith import traceable  # names agent steps in LangSmith (does nothing when tracing is off)

import ado_client
import approval_ledger
from prompt_guard import UNTRUSTED_RULE, check_outgoing, screen, warning_banner, wrap

APP_DIR = Path(__file__).resolve().parent

# ---------------- Chat model (Setup section of the notebook) ----------------
LLM_PROVIDER = os.getenv("LLM_PROVIDER", "ollama").strip().lower()
if LLM_PROVIDER not in ("openai", "ollama"):
    raise ValueError(f"LLM_PROVIDER must be 'openai' or 'ollama', not {LLM_PROVIDER!r}")
OPENAI_MODEL = os.getenv("OPENAI_MODEL", "gpt-6-luna")
OLLAMA_MODEL = os.getenv("OLLAMA_MODEL", "gemma4:e4b")
LLM_MODEL = OPENAI_MODEL if LLM_PROVIDER == "openai" else OLLAMA_MODEL
EMBED_MODEL = os.getenv("EMBED_MODEL", "nomic-embed-text")

APPROVER = os.getenv("APPROVER_NAME", "Toni Kanaan")  # signs the emails and is recorded in the approval ledger
ADO_LIVE = ado_client.AzureDevOpsClient().live  # the gates use it: auto mode never approves live work items

llm = None  # created on first use, so the module can be imported without a key


def make_llm():
    """The chat model selected by LLM_PROVIDER: GPT Luna on OpenAI, or the local Ollama model."""
    if LLM_PROVIDER == "openai":
        from langchain_openai import ChatOpenAI

        key = os.getenv("OPEN_AI_KEY") or os.getenv("OPENAI_API_KEY")
        if not key:
            raise RuntimeError("LLM_PROVIDER=openai needs OPEN_AI_KEY in .env.")
        # GPT Luna only accepts its default temperature, so none is set.
        return ChatOpenAI(model=LLM_MODEL, api_key=key)

    from langchain_ollama import ChatOllama

    return ChatOllama(model=LLM_MODEL, temperature=0)


def get_llm():
    global llm
    if llm is None:
        llm = make_llm()
    return llm


def ask_structured(schema, system_prompt: str, user_prompt: str, retries: int = 2):
    """Call the LLM and force the answer into a Pydantic schema. On failure, retry with the error message."""
    # The run name shows up in LangSmith, e.g. "DeliveryPlan" or "GapAnalysis".
    structured_llm = get_llm().with_structured_output(schema, method="json_schema").with_config(run_name=schema.__name__)
    last_error = None

    for attempt in range(retries + 1):
        try:
            result = structured_llm.invoke(
                # every agent's system prompt ends with the prompt-injection rule (prompt_guard.py)
                [SystemMessage(content=system_prompt + UNTRUSTED_RULE), HumanMessage(content=user_prompt)]
            )
            if result is None:
                raise ValueError("The model returned an empty answer.")
            return result
        except Exception as error:
            last_error = error
            print(f"[retry {attempt + 1}] invalid structured output: {error}")
            user_prompt += (
                f"\n\nYour previous answer was invalid ({error}). "
                "Return JSON that matches the schema exactly."
            )

    raise last_error


def show(title: str, obj) -> None:
    """Pretty-print a heading and a Pydantic object, dict, list or string."""
    print(f"\n========== {title} ==========\n")
    if hasattr(obj, "model_dump"):
        obj = obj.model_dump()
    print(obj if isinstance(obj, str) else json.dumps(obj, indent=2, ensure_ascii=False))



# ====================================================================================================
# 1. Structured schemas and agent prompts (notebook section 1)
# ====================================================================================================

from typing import Literal

from pydantic import BaseModel, Field


class Requirement(BaseModel):
    req_id: str = Field(description="Stable ID such as REQ-001")
    statement: str = Field(description="One clear, testable requirement sentence")
    source_text: str = Field(description="The customer's words this requirement comes from")
    actor: str = Field(description="Who performs or benefits from the requirement")
    business_rules: list[str] = Field(description="Rules, thresholds or conditions that are actually stated")
    exceptions: list[str] = Field(description="Exceptions or edge cases that are actually stated")
    expected_outcome: str = Field(description="What should happen when the requirement is met")


class RequirementExtraction(BaseModel):
    summary: str = Field(description="Two-sentence summary of the request")
    actors: list[str]
    in_scope: list[str]
    out_of_scope: list[str]
    security_constraints: list[str] = Field(description="Security, approval or segregation-of-duties constraints")
    requirements: list[Requirement]


class Classification(BaseModel):
    module: Literal["Finance", "Supply Chain", "Finance and Supply Chain", "Out of scope"]
    sub_areas: list[str] = Field(description="F&O areas such as Accounts payable, Procurement, Inventory")
    work_type: Literal["Functional", "Technical", "Mixed"]
    design_documents: Literal["FDD", "TDD", "FDD and TDD"]
    risk_level: Literal["Low", "Medium", "High"]
    rationale: str


class Gap(BaseModel):
    gap_id: str = Field(description="Stable ID such as GAP-01")
    req_id: str = Field(description="The requirement this gap belongs to")
    missing_information: str
    why_it_matters: str
    clarification_question: str = Field(description="One clear question the customer can answer")
    assumption_if_unanswered: str


class GapAnalysis(BaseModel):
    gaps: list[Gap]
class ClarificationEmail(BaseModel):
    subject: str
    greeting: str = Field(description="Greeting line only, e.g. 'Dear Rana,'")
    introduction: str = Field(description="One-sentence restatement of the request")
    questions: list[str] = Field(description="Clarification questions, one topic each, without numbers")
    reply_by: str = Field(description="Only the requested reply date, e.g. '05 October 2026'")
    closing: str = Field(description="Closing phrase only, e.g. 'Kind regards,', without a name or signature")


class Answer(BaseModel):
    gap_id: str
    answer: str = Field(description="The customer's answer, with their exact values")


class CustomerAnswers(BaseModel):
    answers: list[Answer]
    unanswered_gap_ids: list[str]
    go_live_date: str = Field(default="", description="Go-live date as YYYY-MM-DD if the customer gave one, else empty")
class AcceptanceCriterion(BaseModel):
    ac_id: str = Field(description="ID such as AC-01")
    req_ids: list[str]
    given: str
    when: str
    then: str


class TaskItem(BaseModel):
    task_id: str = Field(description="ID such as T-01")
    kind: Literal["Functional", "Technical"]
    title: str = Field(description="Starts with FUNC: or TECH:")
    description: str
    req_ids: list[str]
    estimate_hours: float


class UserStory(BaseModel):
    story_id: str = Field(description="ID such as US-01")
    title: str = Field(description="Format: [Module] short goal")
    as_a: str
    i_want: str
    so_that: str
    req_ids: list[str]
    acceptance_criteria: list[AcceptanceCriterion]
    tasks: list[TaskItem]


class DeliveryPlan(BaseModel):
    user_stories: list[UserStory]
class DesignSection(BaseModel):
    section_id: str = Field(description="ID such as FDD-1 or TDD-1")
    heading: str
    content: str
    req_ids: list[str] = Field(description="Requirement IDs this section covers")


class DesignDocument(BaseModel):
    doc_type: Literal["FDD", "TDD"]
    title: str
    overview: str
    sections: list[DesignSection]
    open_points: list[str]


class TestCase(BaseModel):
    test_id: str = Field(description="ID such as TC-01")
    test_type: Literal["Functional", "Technical", "Integration", "Regression"]
    title: str
    req_ids: list[str]
    preconditions: str
    steps: list[str]
    expected_result: str


class TestSuite(BaseModel):
    test_cases: list[TestCase]
AgentName = Literal["planning", "functional_design", "technical_design", "tests"]


class ValidationIssue(BaseModel):
    severity: Literal["Blocker", "Major", "Minor"]
    responsible_agent: AgentName
    req_ids: list[str]
    description: str
    fix_instruction: str


class ConsistencyReview(BaseModel):
    issues: list[ValidationIssue]

REQUIREMENTS_SYSTEM = """
You are the Requirements Agent of a Dynamics 365 Finance and Operations (F&O) delivery team.
Turn an informal customer request into structured requirements.

Rules:
- Write one requirement per distinct need and number them REQ-001, REQ-002, ...
- source_text must quote the customer's own words.
- Record only rules, thresholds, exceptions and outcomes that are actually stated.
  Never invent values such as amounts, companies or roles. Missing details are handled later as gaps.
- Put approval, security and segregation-of-duties constraints in security_constraints.
"""

REFINE_SYSTEM = """
You are the Requirements Agent of a Dynamics 365 F&O delivery team.
Update the requirements with the customer's confirmed answers.

Rules:
- Keep the existing REQ IDs; add a new ID only for a genuinely new need the customer confirmed.
- Put confirmed values (amounts, currencies, legal entities, roles, days, dates) into the statements
  and business rules, using the customer's exact values.
- Do not present assumptions as confirmed facts.
"""

CLASSIFY_SYSTEM = """
You are the Gap and Retrieval Agent of a Dynamics 365 F&O delivery team. Classify the request.

- Finance: general ledger, accounts payable, accounts receivable, cash and bank, fixed assets, budgeting.
- Supply Chain: procurement, inventory, warehouse, sales orders, production.
- Functional work = configuration and process only. Technical work = code, integrations, batch jobs, reports.
  Mixed = both.
- design_documents: FDD for functional work, TDD for technical work, "FDD and TDD" for mixed work.
Use the retrieved knowledge as evidence.
"""

GAPS_SYSTEM = """
You are the Gap and Retrieval Agent of a Dynamics 365 F&O delivery team.
Find the missing information that would force a consultant or developer to guess.

Check every requirement for: amounts, thresholds and currency; legal entities in scope; approver roles
and escalation; exceptions; existing versus new data; security and segregation of duties; audit needs;
deadlines.
- Number the gaps GAP-01, GAP-02, ... (at most 8).
- Write each clarification question in plain business language, without X++ or technical jargon.
- Give a safe assumption to use if the customer does not answer.
"""
COMMUNICATION_SYSTEM = """
You are the Communication Agent of a Dynamics 365 F&O delivery team.
Draft a clarification email to the customer.

Follow the company template and communication standard you are given:
- Restate the request in one sentence, then ask the numbered questions (one topic each).
- Plain, polite business language. No technical jargon.
- Ask for a reply date. Never promise a delivery date or a solution.
"""

ANSWERS_SYSTEM = """
You map a customer's reply to open clarification questions (gaps).
- Record an answer only when the reply really answers that gap, and copy the customer's exact values.
- List every gap the reply does not answer in unanswered_gap_ids.
"""
PLANNING_SYSTEM = """
You are the Planning Agent of a Dynamics 365 F&O delivery team.
Create Azure DevOps user stories from the approved specification.

Rules:
- Every requirement ID must be covered by at least one user story and at least one acceptance criterion.
- Story titles use the format "[Module] short goal". Acceptance criteria use Given/When/Then, IDs AC-01, AC-02, ...
- Split the work into tasks, IDs T-01, T-02, ...: functional tasks start with "FUNC:" (configuration,
  FDD, training) and technical tasks start with "TECH:" (X++ extensions, security objects, batch jobs,
  TDD, unit tests). Give each task an estimate in hours.
- The immutable facts are final: use their exact values and never change them.
"""

FDD_SYSTEM = """
You are the Functional Design Agent of a Dynamics 365 F&O delivery team.
Write the Functional Design Document (FDD) for the approved specification.

- Follow the company design document standard: business context, as-is and to-be process,
  configuration and parameters, roles and security, exceptions, reports and audit, assumptions.
- Section IDs FDD-1, FDD-2, ... Each section lists the requirement IDs it covers, and every requirement
  must appear in at least one section.
- Use only the immutable facts and confirmed answers. Cite evidence IDs in the text, e.g. [MS-WF-01].
"""

TDD_SYSTEM = """
You are the Technical Design Agent of a Dynamics 365 F&O delivery team.
Write the Technical Design Document (TDD) for the approved specification.

- Follow the company design and development standards: solution overview, object list (table and form
  extensions, Chain of Command classes, SysOperation batch classes, security privileges and duties),
  logic, data model changes, error handling and logging, performance, deployment.
- Extensions only (no overlayering), CUS prefix, configurable values in a parameter table.
- Section IDs TDD-1, TDD-2, ... Each section lists the requirement IDs it covers, and every requirement
  must appear in at least one section.
- Use only the immutable facts. Cite evidence IDs in the text, e.g. [DES-2025-031].
"""

TEST_SYSTEM = """
You are the Test Agent of a Dynamics 365 F&O delivery team.
Create test cases for the approved specification, IDs TC-01, TC-02, ...

- Types: Functional, Technical, Integration and Regression.
- Every requirement ID needs at least one test case.
- Threshold rules need boundary tests at, just below and just above the limit.
- Security rules need a negative test proving an unauthorized user is blocked.
- Add regression tests for standard processes that must keep working.
- Use the exact values from the immutable facts.
"""

VALIDATION_SYSTEM = """
You are the Validation Agent of a Dynamics 365 F&O delivery team.
Compare the artefacts with the immutable facts and with each other.

Report only real problems:
- a wrong value (amount, currency, legal entity, role, number of days, hold type, deadline),
- a contradiction between the design, the acceptance criteria and the tests,
- behaviour that is not in the approved requirements.
For each issue name the responsible agent: planning, functional_design, technical_design or tests.
If everything is consistent, return an empty list of issues.
"""

REVISION_INSTRUCTIONS = """
Revise your previous output. It failed validation.

Validation feedback:
{feedback}

Lessons from earlier attempts (reflexion memory):
{reflections}

Return the complete corrected artefact. Keep everything that was already correct.
"""


# ====================================================================================================
# 4. Team and owner recommendation (notebook section 4)
# ====================================================================================================

from datetime import date

DEFAULT_TEAM = [
    {"name": "Maya Khoury", "role": "Functional Consultant", "specializations": ["Finance"], "capacity_hours_per_week": 30},
    {"name": "Karim Nassar", "role": "Functional Consultant", "specializations": ["Supply Chain", "Finance"], "capacity_hours_per_week": 30},
    {"name": "Jad Mansour", "role": "Developer", "specializations": ["Finance"], "capacity_hours_per_week": 35},
    {"name": "Lara Aoun", "role": "Developer", "specializations": ["Supply Chain", "Finance"], "capacity_hours_per_week": 35},
    {"name": "Nour Fares", "role": "QA Engineer", "specializations": ["Finance", "Supply Chain"], "capacity_hours_per_week": 30},
]
TEAM_FILE = Path("team.json")  # edit it to match your team; written with the demo team the first time
if not TEAM_FILE.exists():
    TEAM_FILE.write_text(json.dumps(DEFAULT_TEAM, indent=2), encoding="utf-8")
TEAM = json.loads(TEAM_FILE.read_text(encoding="utf-8"))

ROLE_FOR_TASK = {"Functional": "Functional Consultant", "Technical": "Developer"}


def recommend_owners(plan: dict, module: str, team: list, workload: dict, deadline: str) -> dict:
    """Recommend (never assign) an owner per task from role, specialization, open workload and deadline."""
    weeks_left = max((date.fromisoformat(deadline) - date.today()).days / 7, 1)
    load = dict(workload)  # includes hours already recommended in this plan
    recommendations = {}

    for story in plan["user_stories"]:
        for task in story["tasks"]:
            candidates = []
            for member in team:
                if member["role"] != ROLE_FOR_TASK[task["kind"]]:
                    continue
                specialised = any(s in module for s in member["specializations"])
                free_hours = member["capacity_hours_per_week"] * weeks_left - load.get(member["name"], 0)
                candidates.append((specialised, free_hours, member["name"]))

            if not candidates:
                recommendations[task["task_id"]] = {"owner": None, "reason": "No team member has the required role."}
                continue

            specialised, free_hours, name = max(candidates)
            load[name] = load.get(name, 0) + task["estimate_hours"]
            warning = "" if free_hours >= task["estimate_hours"] else " WARNING: not enough capacity before the deadline."
            recommendations[task["task_id"]] = {
                "owner": name,
                "reason": (f"{ROLE_FOR_TASK[task['kind']]}, "
                           f"{'specialised in ' + module if specialised else 'no ' + module + ' specialisation'}, "
                           f"about {free_hours:.0f} free hours before {deadline}.{warning}"),
            }
    return recommendations


# ====================================================================================================
# 5. MCP client: the agents reach knowledge, email, Azure DevOps and SharePoint only through the servers
# ====================================================================================================

import asyncio
import contextlib
import contextvars
import json
import os
import sys
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

from langchain_mcp_adapters.tools import load_mcp_tools
from mcp import ClientSession, StdioServerParameters
from mcp.client.stdio import stdio_client

MCP_LOG = Path("mcp_server.log")  # the servers' stderr ([tool called] lines)


def mcp_server(script: str) -> dict:
    return {"script": str(APP_DIR / script)}  # the server scripts live next to this module


@contextlib.asynccontextmanager
async def mcp_session(server: str):
    """Start a local FastMCP server over stdio and open an MCP session with it.

    The server's stderr goes to mcp_server.log: a notebook's own stderr cannot be handed to a subprocess.
    The whole environment is passed on (MCP only forwards a few safe variables by default), so the servers
    receive the Ollama and Azure DevOps settings.
    """
    params = StdioServerParameters(
        command=sys.executable, args=[MCP_SERVERS[server]["script"]], cwd=str(Path.cwd()), env=dict(os.environ)
    )
    with open(MCP_LOG, "a", encoding="utf-8") as errlog:
        async with stdio_client(params, errlog=errlog) as (read, write):
            async with ClientSession(read, write) as session:
                await session.initialize()
                yield session


def run_async(coro):
    """Run a coroutine from a notebook cell and print the servers' tool log lines.

    It uses a fresh event loop in a worker thread, so it also works in Jupyter on Windows,
    where the notebook's own loop cannot start subprocesses.
    """
    def _run():
        loop = asyncio.ProactorEventLoop() if sys.platform == "win32" else asyncio.new_event_loop()
        try:
            return loop.run_until_complete(coro)
        finally:
            loop.close()

    log_start = MCP_LOG.stat().st_size if MCP_LOG.exists() else 0
    context = contextvars.copy_context()  # keeps MCP tool calls inside the current LangSmith trace
    with ThreadPoolExecutor(max_workers=1) as pool:
        result = pool.submit(context.run, _run).result()

    if MCP_LOG.exists():
        with open(MCP_LOG, encoding="utf-8") as log:
            log.seek(log_start)
            for line in log.read().splitlines():
                if line.startswith("[") and "tool" in line.split("]")[0]:
                    print(line)
    return result


def _to_python(result):
    if isinstance(result, list):  # recent adapters return a list of content blocks
        result = "".join(b.get("text", "") if isinstance(b, dict) else str(b) for b in result)
    try:
        return json.loads(result)
    except (TypeError, json.JSONDecodeError):
        return result


def call_mcp_tools(server: str, calls: list) -> list:
    """Open one MCP session and run several (tool_name, arguments) calls in it."""
    async def _calls():
        async with mcp_session(server) as session:
            tools = {t.name: t for t in await load_mcp_tools(session)}
            results = []
            for name, args in calls:
                try:
                    results.append(_to_python(await tools[name].ainvoke(args)))
                except Exception as error:
                    print(f"[mcp error] {server}.{name}: {error}")
                    results.append({"error": str(error)})
            return results

    return run_async(_calls())


def call_mcp_tool(server: str, tool_name: str, **arguments):
    return call_mcp_tools(server, [(tool_name, arguments)])[0]


def list_mcp_tools(server: str) -> None:
    """tools/list: print the tools a server exposes, with their descriptions and inputs."""
    async def _list():
        async with mcp_session(server) as session:
            return (await session.list_tools()).tools

    print(f"========== TOOLS DISCOVERED ON '{server}' ==========\n")
    for tool in run_async(_list()):
        print(f"Name: {tool.name}")
        print(f"Description: {' '.join(tool.description.split())}")
        print(f"Inputs: {list(tool.inputSchema.get('properties', {}))}\n")


MCP_SERVERS = {
    "knowledge": mcp_server("knowledge_server.py"),
    "email": mcp_server("email_server.py"),
    "devops": mcp_server("devops_server.py"),
    "sharepoint": mcp_server("sharepoint_server.py"),
}


# ---------------- Company standards and email template from the knowledge server ----------------
STANDARD_IDS = ["STD-DEV-01", "STD-DOC-01", "STD-ADO-01", "STD-TEST-01", "STD-COMM-01", "STD-CR-01", "TPL-CLARIFY-01"]
STANDARDS, CLARIFY_TEMPLATE, CR_STANDARD = {}, "", ""


def load_standards() -> None:
    """The agents follow the standards stored in the knowledge base. They are fetched once through get_document."""
    global STANDARDS, CLARIFY_TEMPLATE, CR_STANDARD
    documents = call_mcp_tools("knowledge", [("get_document", {"doc_id": doc_id}) for doc_id in STANDARD_IDS])
    missing = [doc_id for doc_id, doc in zip(STANDARD_IDS, documents) if not isinstance(doc, dict) or "error" in doc]
    if missing:
        raise RuntimeError(f"Standards missing from the knowledge store: {missing}. Run 'python run.py setup' first.")
    STANDARDS = {doc["doc_id"]: doc["text"] for doc in documents}
    CLARIFY_TEMPLATE = STANDARDS["TPL-CLARIFY-01"]
    CR_STANDARD = STANDARDS["STD-CR-01"]


# ====================================================================================================
# 6. Handoff package, delivery agents, validation, document tool and Change Requests (notebook section 6)
# ====================================================================================================

def handoff_package(state: dict) -> str:
    """The structured package passed at every handoff (Coordination section of the planning document)."""
    package = {
        "request_id": state.get("request_id"),
        "source_text": wrap(state.get("source_text", "")),  # untrusted customer text, marked as data
        "requirements": state.get("requirements", []),
        "classification": state.get("classification"),
        "confirmed_answers": state.get("confirmed_answers", {}),
        "assumptions": state.get("assumptions", []),
        "retrieved_evidence": [
            {**e, "snippet": e["snippet"][:400]} for e in state.get("evidence", [])
        ],
        "approval_status": state.get("approval_status", {}),
        "immutable_facts": state.get("locked_facts", []),
    }
    return json.dumps(package, indent=2, ensure_ascii=False)

def plan_summary(plan: dict) -> str:
    lines = []
    for story in plan["user_stories"]:
        lines.append(f"{story['story_id']} {story['title']} ({', '.join(story['req_ids'])})")
        lines += [f"  {ac['ac_id']} ({', '.join(ac['req_ids'])}): Given {ac['given']}, when {ac['when']}, then {ac['then']}"
                  for ac in story["acceptance_criteria"]]
        lines += [f"  {task['task_id']} {task['title']}" for task in story["tasks"]]
    return "\n".join(lines)


@traceable(name="Planning Agent")
def run_planning_agent(state: dict) -> dict:
    user_prompt = f"APPROVED PACKAGE:\n{handoff_package(state)}\n\nWORK ITEM STANDARD:\n{STANDARDS['STD-ADO-01']}"
    return ask_structured(DeliveryPlan, PLANNING_SYSTEM, user_prompt).model_dump()


@traceable(name="Design Agent")
def run_design_agent(kind: str, state: dict) -> dict:
    """Functional (FDD) or Technical (TDD) Design Agent."""
    system_prompt = FDD_SYSTEM if kind == "FDD" else TDD_SYSTEM
    standards = STANDARDS["STD-DOC-01"] if kind == "FDD" else f"{STANDARDS['STD-DOC-01']}\n{STANDARDS['STD-DEV-01']}"
    user_prompt = (
        f"APPROVED PACKAGE:\n{handoff_package(state)}\n\n"
        f"DELIVERY PLAN:\n{plan_summary(state['plan'])}\n\n"
        f"COMPANY STANDARDS:\n{standards}\n\nWrite the {kind}."
    )
    document = ask_structured(DesignDocument, system_prompt, user_prompt).model_dump()
    document["doc_type"] = kind
    return document


@traceable(name="Test Agent")
def run_test_agent(state: dict) -> dict:
    design_sections = [
        f"{s['section_id']} {s['heading']} ({', '.join(s['req_ids'])})"
        for key in ("fdd", "tdd") if state.get(key) for s in state[key]["sections"]
    ]
    user_prompt = (
        f"APPROVED PACKAGE:\n{handoff_package(state)}\n\n"
        f"ACCEPTANCE CRITERIA AND TASKS:\n{plan_summary(state['plan'])}\n\n"
        "DESIGN SECTIONS:\n" + "\n".join(design_sections) + "\n\n"
        f"TEST STANDARD:\n{STANDARDS['STD-TEST-01']}"
    )
    return ask_structured(TestSuite, TEST_SYSTEM, user_prompt).model_dump()

COLUMN_OWNER = {
    "acceptance_criteria": "planning",
    "tasks": "planning",
    "fdd_sections": "functional_design",
    "tdd_sections": "technical_design",
    "tests": "tests",
}


def build_traceability(requirements: list, plan: dict, fdd: dict | None, tdd: dict | None, tests: dict):
    """Map every requirement ID to the acceptance criteria, tasks, design sections and tests that cover it."""
    matrix = {r["req_id"]: {column: [] for column in COLUMN_OWNER} for r in requirements}
    unknown_refs = []

    def add(req_ids, column, ref):
        for req_id in req_ids:
            if req_id in matrix:
                matrix[req_id][column].append(ref)
            else:
                unknown_refs.append((req_id, column, ref))

    for story in plan["user_stories"]:
        for ac in story["acceptance_criteria"]:
            add(ac["req_ids"], "acceptance_criteria", ac["ac_id"])
        for task in story["tasks"]:
            add(task["req_ids"], "tasks", task["task_id"])
    for design, column in ((fdd, "fdd_sections"), (tdd, "tdd_sections")):
        for section in (design or {}).get("sections", []):
            add(section["req_ids"], column, section["section_id"])
    for test in tests["test_cases"]:
        add(test["req_ids"], "tests", test["test_id"])

    return matrix, unknown_refs


def traceability_issues(matrix: dict, unknown_refs: list, design_documents: str) -> list:
    issues = []
    design_owner = "functional_design" if "FDD" in design_documents else "technical_design"

    for req_id, row in matrix.items():
        if not row["acceptance_criteria"]:
            issues.append({"severity": "Major", "responsible_agent": "planning", "req_ids": [req_id],
                           "description": f"{req_id} has no acceptance criterion.",
                           "fix_instruction": f"Add a Given/When/Then acceptance criterion for {req_id}."})
        if not row["fdd_sections"] and not row["tdd_sections"]:
            issues.append({"severity": "Major", "responsible_agent": design_owner, "req_ids": [req_id],
                           "description": f"{req_id} is not covered by any design section.",
                           "fix_instruction": f"Add or extend a design section that covers {req_id}."})
        if not row["tests"]:
            issues.append({"severity": "Major", "responsible_agent": "tests", "req_ids": [req_id],
                           "description": f"{req_id} has no test case.",
                           "fix_instruction": f"Add at least one test case for {req_id}."})

    for req_id, column, ref in unknown_refs:
        issues.append({"severity": "Major", "responsible_agent": COLUMN_OWNER[column], "req_ids": [req_id],
                       "description": f"{ref} references {req_id}, which is not an approved requirement.",
                       "fix_instruction": f"Only reference approved requirement IDs: {', '.join(matrix)}."})
    return issues


def print_matrix(matrix: dict) -> None:
    print(f"{'Requirement':<12}{'Accept. criteria':<22}{'FDD':<16}{'TDD':<16}{'Tests'}")
    for req_id, row in matrix.items():
        print(f"{req_id:<12}{', '.join(row['acceptance_criteria']) or '-':<22}"
              f"{', '.join(row['fdd_sections']) or '-':<16}{', '.join(row['tdd_sections']) or '-':<16}"
              f"{', '.join(row['tests']) or 'MISSING'}")

def artefact_digest(state: dict) -> str:
    """Compact view of all artefacts for the LLM consistency review."""
    lines = ["IMMUTABLE FACTS:", *state["locked_facts"], "", "ACCEPTANCE CRITERIA:"]
    for story in state["plan"]["user_stories"]:
        for ac in story["acceptance_criteria"]:
            lines.append(f"{ac['ac_id']} ({', '.join(ac['req_ids'])}): Given {ac['given']}, "
                         f"when {ac['when']}, then {ac['then']}")
    for key in ("fdd", "tdd"):
        design = state.get(key)
        if design:
            lines += ["", f"{design['doc_type']} SECTIONS:"]
            lines += [f"{s['section_id']} {s['heading']} ({', '.join(s['req_ids'])}): {s['content'][:600]}"
                      for s in design["sections"]]
    lines += ["", "TEST CASES:"]
    lines += [f"{t['test_id']} [{t['test_type']}] ({', '.join(t['req_ids'])}) {t['title']}: "
              f"expected {t['expected_result']}" for t in state["tests"]["test_cases"]]
    return "\n".join(lines)


@traceable(name="Validation Agent")
def validate_package(state: dict):
    """Validation Agent: deterministic traceability checks + LLM consistency review."""
    matrix, unknown_refs = build_traceability(
        state["requirements"], state["plan"], state.get("fdd"), state.get("tdd"), state["tests"]
    )
    issues = traceability_issues(matrix, unknown_refs, state["classification"]["design_documents"])

    review = ask_structured(ConsistencyReview, VALIDATION_SYSTEM, artefact_digest(state))
    issues += [issue.model_dump() for issue in review.issues]

    for number, issue in enumerate(issues, start=1):
        issue["issue_id"] = f"V-{number:02d}"
    return matrix, issues


AGENT_SPECS = {
    # agent name: (system prompt, output schema, state key)
    "planning": (PLANNING_SYSTEM, DeliveryPlan, "plan"),
    "functional_design": (FDD_SYSTEM, DesignDocument, "fdd"),
    "technical_design": (TDD_SYSTEM, DesignDocument, "tdd"),
    "tests": (TEST_SYSTEM, TestSuite, "tests"),
}


@traceable(name="Reflexion revision")
def revise_artefact(agent: str, state: dict, issues: list) -> dict:
    """Send the artefact back to the responsible agent with the validation feedback and reflexion memory."""
    system_prompt, schema, key = AGENT_SPECS[agent]
    feedback = "\n".join(
        f"- {i['issue_id']} [{i['severity']}] {i['description']} Fix: {i['fix_instruction']}"
        for i in issues if i["responsible_agent"] == agent
    )
    user_prompt = (
        f"APPROVED PACKAGE:\n{handoff_package(state)}\n\n"
        f"YOUR PREVIOUS OUTPUT:\n{json.dumps(state[key], indent=2, ensure_ascii=False)}\n\n"
        + REVISION_INSTRUCTIONS.format(feedback=feedback, reflections="\n".join(state.get("reflections", [])) or "none")
    )
    return ask_structured(schema, system_prompt, user_prompt).model_dump()


def is_blocking(issue: dict) -> bool:
    return issue["severity"] in ("Blocker", "Major")

from datetime import date

import docx
from docx.shared import Pt


def _new_document(title: str, request_id: str, subtitle: str):
    document = docx.Document()
    document.styles["Normal"].font.size = Pt(10.5)
    document.add_heading(title, 0)
    document.add_paragraph(f"Request: {request_id}    |    {subtitle}    |    Generated: {date.today()}")
    return document


def export_design_document(design: dict, request_id: str, locked_facts: list, folder: Path) -> Path:
    """Document tool: write an FDD or TDD as a Word file."""
    document = _new_document(design["title"], request_id, design["doc_type"])
    document.add_heading("Overview", 1)
    document.add_paragraph(design["overview"])
    document.add_heading("Approved facts (immutable)", 1)
    for fact in locked_facts:
        document.add_paragraph(fact, style="List Bullet")
    for section in design["sections"]:
        document.add_heading(f"{section['section_id']}  {section['heading']}", 1)
        document.add_paragraph(section["content"])
        document.add_paragraph().add_run(f"Covers: {', '.join(section['req_ids'])}").italic = True
    if design["open_points"]:
        document.add_heading("Open points", 1)
        for point in design["open_points"]:
            document.add_paragraph(point, style="List Bullet")

    path = folder / f"{request_id}_{design['doc_type']}.docx"
    document.save(path)
    return path


def export_test_document(tests: dict, matrix: dict, request_id: str, folder: Path) -> Path:
    """Document tool: write the test cases and the traceability matrix as a Word file."""
    document = _new_document("Test Cases and Traceability", request_id, "Test plan")

    document.add_heading("Traceability matrix", 1)
    columns = ["acceptance_criteria", "fdd_sections", "tdd_sections", "tests"]
    table = document.add_table(rows=1, cols=len(columns) + 1)
    table.style = "Light Grid Accent 1"
    for cell, text in zip(table.rows[0].cells, ["Requirement", "Acceptance criteria", "FDD", "TDD", "Tests"]):
        cell.text = text
    for req_id, row in matrix.items():
        cells = table.add_row().cells
        cells[0].text = req_id
        for cell, column in zip(cells[1:], columns):
            cell.text = ", ".join(row[column]) or "-"

    document.add_heading("Test cases", 1)
    for test in tests["test_cases"]:
        document.add_heading(f"{test['test_id']}  {test['title']}", 2)
        document.add_paragraph(f"Type: {test['test_type']}    Requirements: {', '.join(test['req_ids'])}")
        document.add_paragraph(f"Preconditions: {test['preconditions']}")
        for step in test["steps"]:
            document.add_paragraph(step, style="List Number")
        document.add_paragraph(f"Expected result: {test['expected_result']}")

    path = folder / f"{request_id}_Test_Cases.docx"
    document.save(path)
    return path

class CRAssessment(BaseModel):
    request_type: Literal["Support", "Change request", "Support and change request"]
    cr_required: bool = Field(description="True when any part of the request needs a Change Request")
    reasons: list[str] = Field(description="Why, each reason naming the rule of the CR standard that applies")
    change_items: list[str] = Field(description="Parts that change the agreed solution (empty when none)")
    support_items: list[str] = Field(description="Parts that are support: defects, data or setup fixes, questions")


class ChangeRequestContent(BaseModel):
    change_name: str = Field(description="Short name of the change, e.g. 'Vendor invoice approval above 10,000 USD'")
    business_requirement: str = Field(description="What the customer needs and why, in business language (1-3 short paragraphs)")
    proposed_solution: str = Field(description="How it will be delivered in Dynamics 365 F&O, in business language")
    customization_required: bool
    technical_complexity: Literal["Low", "Medium", "High"]
    business_priority: Literal["Low", "Medium", "High"] = Field(description="From the customer's urgency and impact")
    recommendation: str = Field(description="Recommendation and best practice: standard features first, customization only where needed")
    affected_areas: list[str] = Field(description="Forms, journals, workflows, reports, entities or processes affected")
    assumptions: list[str]


CR_ASSESSMENT_SYSTEM = """
You are the Change Request Assessment Agent of a Dynamics 365 F&O partner.
Decide whether a customer request needs a Change Request (CR), using only the company CR standard you are given.
- A defect in functionality that already works as agreed, a data or setup problem, a how-to question or an
  access issue is support. New or changed functionality, customizations, reports, integrations, workflows,
  fields, processes or scope are changes.
- Name the rule of the standard that applies in each reason.
- When the request mixes both, list each part and set cr_required to true.
"""

CR_WRITER_SYSTEM = """
You are the Change Request Agent of a Dynamics 365 F&O implementation partner.
Write the partner's parts of the ERP Change/Additional Request Form for the customer to review and sign.
- Business language for the customer's managers; technical object names only in affected_areas.
- Use only facts from the request, the confirmed facts and the plan. Never invent amounts, dates or names.
- Recommend standard features first and say clearly when customization is required and why.
- Do not mention prices. The estimated effort is added from the plan.
"""


@traceable(name="CR Assessment Agent")
def assess_change_request(request_text: str, context: str = "") -> CRAssessment:
    user_prompt = f"Company CR standard:\n{CR_STANDARD}\n\nCustomer request:\n{wrap(request_text)}"
    if context:
        user_prompt += f"\n\nContext:\n{context}"
    return ask_structured(CRAssessment, CR_ASSESSMENT_SYSTEM, user_prompt)


@traceable(name="Change Request Agent")
def write_change_request(request_text: str, change_items: list, context: str = "") -> ChangeRequestContent:
    user_prompt = (f"Company CR standard:\n{CR_STANDARD}\n\nCustomer request:\n{wrap(request_text)}\n\n"
                   "Parts that are changes:\n" + "\n".join(f"- {item}" for item in change_items))
    if context:
        user_prompt += f"\n\nConfirmed facts and plan:\n{context}"
    return ask_structured(ChangeRequestContent, CR_WRITER_SYSTEM, user_prompt)


import math
import re
from datetime import date, datetime

import docx
from docx.shared import Pt

CR_PROJECT_NAME = os.getenv("CR_PROJECT_NAME", "Dynamics 365 F&O")
CR_SIGN_OFF = ["Functional Lead - Customer", "Head of Department - Customer", "IT Manager - Customer",
               "Project Manager - Partner"]


def form_date(received: str = "") -> str:
    """'2026-09-25T09:14:00Z' -> '25/09/2026', the form's date format (today when the date is unknown)."""
    try:
        return datetime.fromisoformat(received.replace("Z", "+00:00")).strftime("%d/%m/%Y")
    except ValueError:
        return date.today().strftime("%d/%m/%Y")


def cr_file_name(change_number) -> str:
    number = str(change_number)
    return re.sub(r"[^A-Za-z0-9_-]", "_", number if not number.isdigit() else f"CR-{number}") + "_Change_Request.docx"


def export_change_request(cr: dict, meta: dict, folder: Path) -> Path:
    """Document tool: the ERP Change/Additional Request Form (the company layout) as a Word file.
    meta: change_number, requested_by, date_of_request, phase, effort_hours."""
    document = docx.Document()
    document.styles["Normal"].font.name = "Arial"
    document.styles["Normal"].font.size = Pt(10)
    document.add_paragraph("ERP Change/Additional Request Form", style="Title")

    header = document.add_table(rows=4, cols=4)
    header.style = "Table Grid"
    values = [("Project Name", CR_PROJECT_NAME, "Change Number", meta["change_number"]),
              ("Requested By", meta["requested_by"], "Date of Request", meta["date_of_request"]),
              ("Required Phase of Project", meta["phase"], "Business Priority (High, Medium, Low)", cr["business_priority"]),
              ("Change Name", cr["change_name"], "Technical Complexity (High, Medium, Low)", cr["technical_complexity"])]
    for row, texts in zip(header.rows, values):
        for column, (cell, text) in enumerate(zip(row.cells, texts)):
            cell.paragraphs[0].add_run(str(text)).bold = column % 2 == 0  # labels in bold

    def section(title: str, blocks=()) -> None:
        """One form section: a bold title row and a body row. blocks: (kind, text) with kind heading/text/bullet."""
        document.add_paragraph()
        table = document.add_table(rows=2, cols=1)
        table.style = "Table Grid"
        table.rows[0].cells[0].paragraphs[0].add_run(title).bold = True
        body = table.rows[1].cells[0]
        for number, (kind, text) in enumerate(blocks):
            paragraph = body.paragraphs[0] if number == 0 else body.add_paragraph()
            run = paragraph.add_run(f"\u2022  {text}" if kind == "bullet" else text)
            run.bold = kind == "heading"

    def paragraphs(text: str) -> list:
        return [("text", line.strip()) for line in text.split("\n") if line.strip()]

    section("Change Description: (To be filled by the Partner)",
            [("heading", "Business Requirement"), *paragraphs(cr["business_requirement"]),
             ("heading", "Proposed Solution"), *paragraphs(cr["proposed_solution"])])
    section("Business Justification (Value): (Risk, productivity, efficiency, financial uplift) "
            "(To be filled by Customer Business)")
    section("Impact of not implementing the Change: (To be filled by Customer Business)")
    effort = meta["effort_hours"]
    section("Recommendation/Best Practice: (To be filled by the Partner)",
            [("heading", f"Customization Required: {'Yes' if cr['customization_required'] else 'No'}"),
             *paragraphs(cr["recommendation"]),
             ("heading", "Affected Areas"), *[("bullet", area) for area in cr["affected_areas"]],
             ("heading", "Estimated Effort"),
             ("text", f"{effort:g} hours (about {math.ceil(effort / 8)} working days), from the delivery plan"),
             *([("heading", "Assumptions")] + [("bullet", a) for a in cr["assumptions"]] if cr["assumptions"] else [])])

    document.add_paragraph()
    approval = document.add_table(rows=1 + len(CR_SIGN_OFF), cols=1)
    approval.style = "Table Grid"
    approval.rows[0].cells[0].paragraphs[0].add_run("For Approval:").bold = True
    for row, role in zip(approval.rows[1:], CR_SIGN_OFF):
        row.cells[0].text = f"Sign off by: ______________________ ({role})"

    folder.mkdir(parents=True, exist_ok=True)
    path = folder / cr_file_name(meta["change_number"])
    document.save(path)
    return path


CR_PARAGRAPH = {
    "Change request": "This request changes the agreed solution, so it needs a Change Request.",
    "Support and change request": "Part of this request changes the agreed solution, so that part needs a Change Request.",
}


def cr_reply_paragraph(cr: dict | None) -> str:
    """The acknowledgement paragraph that introduces the attached Change Request ("" when none is needed)."""
    if not cr:
        return ""
    return (f"**Change Request:**\n{CR_PARAGRAPH.get(cr['request_type'], CR_PARAGRAPH['Change request'])} "
            f"Please review the attached Change Request **{cr['number']}**, complete the business justification "
            "and impact sections, and return it signed. Work on the change starts once it is approved.\n\n")

import re
def print_plan(plan: dict, owners: dict | None = None) -> None:
    for story in plan["user_stories"]:
        print(f"\n{story['story_id']}  {story['title']}   covers {', '.join(story['req_ids'])}")
        print(f"   As a {story['as_a']}, I want {story['i_want']} so that {story['so_that']}.")
        for ac in story["acceptance_criteria"]:
            print(f"   {ac['ac_id']} ({', '.join(ac['req_ids'])}): Given {ac['given']}, when {ac['when']}, then {ac['then']}")
        for task in story["tasks"]:
            owner = f"  -> recommended: {owners[task['task_id']]['owner']}" if owners and task["task_id"] in owners else ""
            print(f"   - {task['task_id']} {task['title']} [{task['estimate_hours']}h]{owner}")
REPLY_BY = (date.today() + timedelta(days=4)).strftime("%d %B %Y")
DECISIONS = {"approve": "approved", "revise": "revise", "reject": "rejected"}


def audit(agent: str, message: str) -> list:
    line = f"[{agent}] {message}"
    print(line)
    return [line]


def guarded(state: dict, request: dict, outgoing: str | None = None) -> dict:
    """A human gate with the prompt-injection findings on top: the reviewer sees them first, and auto mode never
    approves a gate that has findings. outgoing: an email body to check before it leaves (output check)."""
    findings = list(state.get("security_findings", []))
    if outgoing is not None:
        email = state.get("email") or {}
        source = state.get("source_text") or email.get("body", "")
        known = (state.get("requester_address", ""), email.get("from_address", ""), *email.get("cc", []))
        findings += [{**f, "source": "outgoing email"} for f in check_outgoing(outgoing, source, known)]
    if not findings:
        return request
    return {**request, "security": findings, "artefact": warning_banner(findings) + request["artefact"]}


def format_evidence(evidence: list) -> str:
    return "\n".join(f"[{e['doc_id']}] {e['snippet'][:400]}" for e in evidence) or "none"


def render_email(email: ClarificationEmail) -> str:
    questions = "\n".join(  # the numbering is added here, so any number the model wrote is removed
        f"{number}. {re.sub(r'^\s*\d+[.)]\s*', '', question)}"
        for number, question in enumerate(email.questions, start=1)
    )
    return (f"{email.greeting}\n\n{email.introduction}\n\n{questions}\n\n"
            f"Could you please reply by {email.reply_by}?\n\n{email.closing}\n{APPROVER}\n"
            "Dynamics 365 F&O Delivery Team")


def captured(function, *args) -> str:
    """Return what a print-based helper (print_plan, print_matrix) would print."""
    buffer = io.StringIO()
    with contextlib.redirect_stdout(buffer):
        function(*args)
    return buffer.getvalue()


# ====================================================================================================
# 7. The supervisor workflow (notebook section 7)
# ====================================================================================================

MAX_REVISIONS = 3


def lock_once(current: list, new: list) -> list:
    """Immutable facts: once set, later updates are ignored."""
    return current if current else new


def merge_dicts(current: dict, new: dict) -> dict:
    return {**current, **new}


class DeliveryState(TypedDict, total=False):
    request_id: str
    message_id: str
    source_text: str
    requested_by: str
    requester_address: str
    request_subject: str
    sharepoint_folder: dict
    request_date: str
    cr_assessment: dict
    extraction: dict
    requirements: list
    classification: dict
    evidence: list
    gaps: list
    clarification_email: dict
    clarification_draft: dict
    email_feedback: str
    confirmed_answers: dict
    assumptions: list
    deadline: str
    spec_feedback: str
    locked_facts: Annotated[list, lock_once]
    approval_status: Annotated[dict, merge_dicts]
    plan: dict
    owner_recommendations: dict
    fdd: dict | None
    tdd: dict | None
    tests: dict
    matrix: dict
    issues: list
    revision_target: str | None
    revision_count: int
    reflections: Annotated[list, operator.add]
    security_findings: Annotated[list, operator.add]  # prompt-injection findings, shown at every gate
    audit_log: Annotated[list, operator.add]
    status: str
    exported_files: list
    work_items: list

def intake(state: DeliveryState) -> dict:
    if state.get("message_id"):
        email = call_mcp_tool("email", "get_email", message_id=state["message_id"])
        if "error" in email:
            raise RuntimeError(email["error"])
        text, source = email["body"], f"support email '{email['subject']}' from {email['from_address']}"
        sender = {"requested_by": email["from_name"], "request_date": form_date(email["received"]),
                  "requester_address": email["from_address"], "request_subject": email["subject"]}
    else:
        text, source, sender = state["source_text"], "manual input", {}
    text, findings = screen(text)  # input check: hidden text removed, instruction-like text flagged
    log = audit("supervisor", f"{state['request_id']} received via {source}")
    if findings:
        log += audit("guard", f"possible prompt injection in the request: {', '.join(f['rule'] for f in findings)}")
    return {"source_text": text, **sender, "security_findings": [{**f, "source": "request email"} for f in findings],
            "audit_log": log}


def requirements_agent(state: DeliveryState) -> dict:
    extraction = ask_structured(RequirementExtraction, REQUIREMENTS_SYSTEM, f"Customer request:\n{wrap(state['source_text'])}")
    requirements = [r.model_dump() for r in extraction.requirements]
    return {"extraction": extraction.model_dump(), "requirements": requirements,
            "audit_log": audit("requirements", f"{len(requirements)} requirements extracted")}


def gap_retrieval_agent(state: DeliveryState) -> dict:
    extraction = state["extraction"]

    broad = call_mcp_tool("knowledge", "search_fo_knowledge", query=extraction["summary"], k=4)
    broad = broad if isinstance(broad, list) else []
    classification = ask_structured(
        Classification, CLASSIFY_SYSTEM,
        f"Requirements:\n{json.dumps(extraction, indent=2)}\n\nRetrieved knowledge:\n{format_evidence(broad)}",
    ).model_dump()

    per_requirement = call_mcp_tools("knowledge", [
        ("search_fo_knowledge", {"query": r["statement"], "module": classification["module"], "k": 3})
        for r in state["requirements"]
    ])
    evidence = {}
    for hit in broad + [h for hits in per_requirement if isinstance(hits, list) for h in hits]:
        if hit["source_type"] != "email_template":
            evidence.setdefault(hit["doc_id"], hit)
    evidence = list(evidence.values())

    gaps = ask_structured(GapAnalysis, GAPS_SYSTEM, (
        f"Original request:\n{wrap(state['source_text'])}\n\n"
        f"Requirements:\n{json.dumps(state['requirements'], indent=2)}\n\n"
        f"Classification:\n{json.dumps(classification, indent=2)}\n\n"
        f"Retrieved knowledge:\n{format_evidence(evidence)}"
    ))
    gaps = [g.model_dump() for g in gaps.gaps]

    log = audit("gap_retrieval", f"{classification['module']} / {classification['work_type']} -> "
                                 f"{classification['design_documents']}; evidence {[e['doc_id'] for e in evidence]}; "
                                 f"{len(gaps)} gaps")
    return {"classification": classification, "evidence": evidence, "gaps": gaps, "audit_log": log}


def cr_assessment_agent(state: DeliveryState) -> dict:
    assessment = assess_change_request(state["source_text"],
                                       context=f"Classification:\n{json.dumps(state['classification'], indent=2)}")
    verdict = "CR required" if assessment.cr_required else "no CR needed"
    return {"cr_assessment": assessment.model_dump(),
            "audit_log": audit("cr_assessment", f"{assessment.request_type}: {verdict}")}


def route_after_gaps(state: DeliveryState) -> str:
    return "communication_agent" if state["gaps"] else "spec_approval"

def communication_agent(state: DeliveryState) -> dict:
    questions = "\n".join(f"{g['gap_id']}: {g['clarification_question']}" for g in state["gaps"])
    user_prompt = (
        f"Company template:\n{CLARIFY_TEMPLATE}\n\nCommunication standard:\n{STANDARDS['STD-COMM-01']}\n\n"
        f"Customer request:\n{wrap(state['source_text'])}\n\nOpen questions (keep their order):\n{questions}\n\n"
        f"Reply date: {REPLY_BY}\nConsultant name: {APPROVER}"
    )
    if state.get("email_feedback"):
        user_prompt += f"\n\nReviewer feedback on the previous draft: {state['email_feedback']}"

    email = ask_structured(ClarificationEmail, COMMUNICATION_SYSTEM, user_prompt)
    return {"clarification_email": {"subject": email.subject, "body": render_email(email)},
            "audit_log": audit("communication", f"clarification email drafted ({len(email.questions)} questions)")}


def approve_clarification(state: DeliveryState) -> Command:
    email = state["clarification_email"]
    review = interrupt(guarded(state, {"type": "approval", "title": "Clarification email to the customer",
                                       "artefact": f"Subject: {email['subject']}\n\n{email['body']}"},
                               outgoing=email["body"]))

    decision = DECISIONS[review["decision"]]
    record = approval_ledger.record_decision("customer_email", state.get("message_id") or state["request_id"],
                                             APPROVER, decision, review["comments"])
    log = audit("human", f"clarification email {decision} ({record['approval_id']})")

    if decision == "revise":
        return Command(goto="communication_agent", update={"email_feedback": review["comments"], "audit_log": log})
    if decision == "rejected":
        return Command(goto=END, update={"status": "stopped: clarification email rejected", "audit_log": log})

    update = {"approval_status": {"clarification_email": record["approval_id"]}}
    if not state.get("message_id"):  # manual input: there is no email to reply to
        return Command(goto="wait_for_customer", update={**update, "audit_log": log})

    draft = call_mcp_tool("email", "save_reply_draft", message_id=state["message_id"], subject=email["subject"],
                          body=email["body"], approval_id=record["approval_id"])
    if "error" in draft:
        raise RuntimeError(draft["error"])
    log += audit("email", f"reply draft saved ({draft['mode']}, {draft['status']})")
    return Command(goto="send_clarification", update={**update, "clarification_draft": draft, "audit_log": log})


def send_gate(draft: dict, label: str, state: dict) -> dict:
    """Second human decision: send this exact draft. Returns the state update."""
    review = interrupt(guarded(state, {
        "type": "approval", "title": f"Send the {label} to {', '.join(draft['to'] + draft['cc'])}?",
        "artefact": f"Subject: {draft['subject']}\n\n{draft['body']}",
        "options": ["approve", "reject"], "live": draft["mode"] == "live"}, outgoing=draft["body"]))
    decision = DECISIONS[review["decision"]]
    record = approval_ledger.record_decision("send_email", draft["draft_id"], APPROVER, decision, review["comments"])
    log = audit("human", f"sending the {label} {decision} ({record['approval_id']})")
    if decision != "approved":
        return {"audit_log": log + audit("email", "not sent: the draft stays in the mailbox")}

    sent = call_mcp_tool("email", "send_draft", draft_id=draft["draft_id"], approval_id=record["approval_id"])
    if "error" in sent:
        raise RuntimeError(sent["error"])
    return {"audit_log": log + audit("email", f"{label} sent to {', '.join(sent['to'])} ({sent['status']})")}


def send_clarification(state: DeliveryState) -> dict:
    return send_gate(state["clarification_draft"], "clarification email", state)


def wait_for_customer(state: DeliveryState) -> dict:
    reply = interrupt({"type": "customer_reply", "title": "Waiting for the customer's answers",
                       "artefact": state["clarification_email"]["body"]})["reply"]

    reply, findings = screen(reply)  # the customer's reply is untrusted too
    gap_list = "\n".join(f"{g['gap_id']}: {g['clarification_question']}" for g in state["gaps"])
    mapped = ask_structured(CustomerAnswers, ANSWERS_SYSTEM, f"Open gaps:\n{gap_list}\n\nCustomer reply:\n{wrap(reply)}")

    gap_ids = {g["gap_id"] for g in state["gaps"]}
    confirmed = {a.gap_id: a.answer for a in mapped.answers if a.gap_id in gap_ids}
    assumptions = [f"{g['gap_id']}: {g['assumption_if_unanswered']}" for g in state["gaps"] if g["gap_id"] not in confirmed]
    try:
        deadline = date.fromisoformat(mapped.go_live_date).isoformat()
    except ValueError:
        deadline = (date.today() + timedelta(days=90)).isoformat()

    log = audit("communication", f"{len(confirmed)} answers confirmed, {len(assumptions)} assumptions, "
                                 f"deadline {deadline}")
    if findings:
        log += audit("guard", f"possible prompt injection in the customer's reply: {', '.join(f['rule'] for f in findings)}")
    return {"confirmed_answers": confirmed, "assumptions": assumptions, "deadline": deadline,
            "security_findings": [{**f, "source": "customer reply"} for f in findings], "audit_log": log}

def refine_specification(state: DeliveryState) -> dict:
    answers = "\n".join(f"{gap_id}: {answer}" for gap_id, answer in state.get("confirmed_answers", {}).items())
    user_prompt = (
        f"Current requirements:\n{json.dumps(state['requirements'], indent=2)}\n\n"
        f"Confirmed answers (the customer's words):\n{wrap(answers) if answers else 'none'}\n\n"
        f"Assumptions (not confirmed):\n" + ("\n".join(state.get("assumptions", [])) or "none")
    )
    if state.get("spec_feedback"):
        user_prompt += f"\n\nReviewer feedback on the previous specification: {state['spec_feedback']}"

    refined = ask_structured(RequirementExtraction, REFINE_SYSTEM, user_prompt)
    requirements = [r.model_dump() for r in refined.requirements]
    return {"extraction": refined.model_dump(), "requirements": requirements,
            "audit_log": audit("requirements", f"specification refined: {len(requirements)} requirements")}


def build_locked_facts(state: DeliveryState) -> list:
    facts = [f"Request ID: {state['request_id']}"]
    facts += [f"{r['req_id']}: {r['statement']}" for r in state["requirements"]]
    facts += [f"Customer confirmed ({gap_id}): {answer}" for gap_id, answer in state.get("confirmed_answers", {}).items()]
    facts += [f"Security constraint: {c}" for c in state["extraction"].get("security_constraints", [])]
    facts.append("Cited evidence: " + ", ".join(e["doc_id"] for e in state.get("evidence", [])))
    return facts


def spec_approval(state: DeliveryState) -> Command:
    spec = "\n".join(
        [f"Classification: {state['classification']['module']} / {state['classification']['work_type']} "
         f"-> {state['classification']['design_documents']}", "", "REQUIREMENTS:"]
        + [f"  {r['req_id']}: {r['statement']}" for r in state["requirements"]]
        + ["", "CONFIRMED ANSWERS:"] + [f"  {k}: {v}" for k, v in state.get("confirmed_answers", {}).items()]
        + ["", "ASSUMPTIONS (please check):"] + [f"  {a}" for a in state.get("assumptions", [])]
    )
    review = interrupt(guarded(state, {"type": "approval", "title": "Specification approval", "artefact": spec}))

    decision = DECISIONS[review["decision"]]
    record = approval_ledger.record_decision("specification", state["request_id"], APPROVER, decision, review["comments"])
    log = audit("human", f"specification {decision} ({record['approval_id']})")

    if decision == "revise":
        return Command(goto="refine_specification", update={"spec_feedback": review["comments"], "audit_log": log})
    if decision == "rejected":
        return Command(goto=END, update={"status": "stopped: specification rejected", "audit_log": log})

    facts = build_locked_facts(state)
    log += audit("supervisor", f"{len(facts)} immutable facts locked")
    return Command(goto="planning_agent", update={
        "locked_facts": facts, "approval_status": {"specification": record["approval_id"]}, "audit_log": log,
    })

def planning_agent(state: DeliveryState) -> dict:
    revising = state.get("revision_target") == "planning"
    plan = revise_artefact("planning", state, state["issues"]) if revising else run_planning_agent(state)

    workload = call_mcp_tool("devops", "get_team_workload")
    workload = {} if "error" in workload else workload
    deadline = state.get("deadline") or (date.today() + timedelta(days=90)).isoformat()
    owners = recommend_owners(plan, state["classification"]["module"], TEAM, workload, deadline)

    tasks = sum(len(s["tasks"]) for s in plan["user_stories"])
    return {"plan": plan, "owner_recommendations": owners,
            "audit_log": audit("planning", f"{'revised: ' if revising else ''}{len(plan['user_stories'])} stories, "
                                           f"{tasks} tasks, owners recommended")}


def route_after_planning(state: DeliveryState):
    return "validation_agent" if state.get("revision_target") else ["functional_design", "technical_design"]


def design_node(kind: str):
    key = kind.lower()
    agent = "functional_design" if kind == "FDD" else "technical_design"

    def node(state: DeliveryState) -> dict:
        if kind not in state["classification"]["design_documents"]:
            return {key: None, "audit_log": audit(agent, f"{kind} not required")}
        revising = state.get("revision_target") == agent
        document = revise_artefact(agent, state, state["issues"]) if revising else run_design_agent(kind, state)
        document["doc_type"] = kind
        return {key: document, "audit_log": audit(agent, f"{'revised ' if revising else ''}{kind}: "
                                                         f"{len(document['sections'])} sections")}

    return node


def route_after_design(state: DeliveryState) -> str:
    return "validation_agent" if state.get("revision_target") else "test_agent"


def test_agent(state: DeliveryState) -> dict:
    revising = state.get("revision_target") == "tests"
    tests = revise_artefact("tests", state, state["issues"]) if revising else run_test_agent(state)
    return {"tests": tests, "audit_log": audit("tests", f"{'revised: ' if revising else ''}"
                                                        f"{len(tests['test_cases'])} test cases")}

AGENT_NODES = {"planning": "planning_agent", "functional_design": "functional_design",
               "technical_design": "technical_design", "tests": "test_agent"}


def validation_agent(state: DeliveryState) -> Command:
    matrix, issues = validate_package(state)
    blocking = [i for i in issues if is_blocking(i)]
    count = state.get("revision_count", 0)
    log = audit("validation", f"round {count}: {len(issues)} issues, {len(blocking)} blocking")

    if blocking and count < MAX_REVISIONS:
        target = min((i["responsible_agent"] for i in blocking), key=list(AGENT_NODES).index)
        docs = state["classification"]["design_documents"]
        if target == "functional_design" and "FDD" not in docs:
            target = "technical_design"
        if target == "technical_design" and "TDD" not in docs:
            target = "functional_design"
        lessons = [f"Round {count + 1}: {i['responsible_agent']} - {i['description']} -> {i['fix_instruction']}"
                   for i in blocking]
        log += audit("supervisor", f"returning the package to {target}")
        return Command(goto=AGENT_NODES[target], update={
            "matrix": matrix, "issues": issues, "revision_target": target,
            "revision_count": count + 1, "reflections": lessons, "audit_log": log,
        })

    return Command(goto="final_approval", update={
        "matrix": matrix, "issues": issues, "revision_target": None, "audit_log": log,
    })


def final_approval(state: DeliveryState) -> Command:
    remaining = "\n".join(f"  {i['issue_id']} [{i['severity']}] {i['description']}" for i in state["issues"]) or "  none"
    cr = state.get("cr_assessment") or {}
    cr_line = f"required ({cr['request_type']}), exported with the documents" if cr.get("cr_required") else "not needed"
    summary = (
        f"USER STORIES, TASKS AND RECOMMENDED OWNERS:{captured(print_plan, state['plan'], state['owner_recommendations'])}\n"
        f"DESIGNS: " + ", ".join(f"{d['doc_type']} ({len(d['sections'])} sections)" for d in (state.get("fdd"), state.get("tdd")) if d)
        + f"\nTESTS: {len(state['tests']['test_cases'])} test cases\nCHANGE REQUEST: {cr_line}\n\n"
        f"TRACEABILITY MATRIX:\n{captured(print_matrix, state['matrix'])}\n"
        f"REMAINING VALIDATION ISSUES:\n{remaining}"
    )
    review = interrupt(guarded(state, {"type": "approval", "title": "Final approval of the delivery package",
                                       "artefact": summary, "live": ADO_LIVE}))

    decision = DECISIONS[review["decision"]]
    record = approval_ledger.record_decision("final_delivery_package", state["request_id"], APPROVER,
                                             decision, review["comments"])
    log = audit("human", f"delivery package {decision} ({record['approval_id']})")

    if decision == "revise":
        feedback = {"severity": "Major", "responsible_agent": "planning", "req_ids": [], "issue_id": "H-01",
                    "description": f"Human reviewer: {review['comments']}", "fix_instruction": review["comments"]}
        return Command(goto="planning_agent", update={
            "issues": [feedback], "revision_target": "planning",
            "reflections": [f"Human reviewer: {review['comments']}"], "audit_log": log,
        })
    if decision == "rejected":
        return Command(goto=END, update={"status": "stopped: delivery package rejected", "audit_log": log})
    return Command(goto="publish", update={"approval_status": {"final_delivery_package": record["approval_id"]},
                                           "audit_log": log})


def customer_for(address: str) -> str:
    """The customer name for SharePoint folders, from the same Feature/project matching as the support flow."""
    if not address:
        return "Unknown customer"
    suggestion = call_mcp_tool("devops", "suggest_features_for_sender", sender_email=address)
    if isinstance(suggestion, dict) and suggestion.get("default"):
        return suggestion["default"]["title"].split(" - ")[0]
    match = (suggestion.get("match") or {}) if isinstance(suggestion, dict) else {}
    return match.get("project") or address.split("@")[-1].split(".")[0].title()


def file_on_sharepoint(state: DeliveryState, files: list, created: list) -> tuple:
    """File the exported documents and link the request folder on the new User Stories."""
    subject = state.get("request_subject") or state["request_id"]
    folder_args = {"customer": customer_for(state.get("requester_address", "")), "item_id": state["request_id"],
                   "description": re.sub(r"^(f&o support|re|fw|fwd)\s*[-:]\s*", "", subject, flags=re.IGNORECASE)}
    by_type = {"Design": [p for p in files if p.stem.endswith(("_FDD", "_TDD"))],
               "Test": [p for p in files if p.stem.endswith("_Test_Cases")],
               "Change Request": [p for p in files if p.stem.endswith("_Change_Request")]}
    calls = [("file_documents", {**folder_args, "doc_type": doc_type, "file_paths": [str(p) for p in paths]})
             for doc_type, paths in by_type.items() if paths]
    *filings, folder = call_mcp_tools("sharepoint", calls + [("ticket_folder_url", folder_args)])
    uploaded = sum(f["status"] == "uploaded" for r in filings if isinstance(r, dict) for f in r.get("files", []))
    stories = [item for item in created if not item.get("parent_id")]
    links = call_mcp_tools("devops", [("add_document_link", {
        "work_item_id": item["id"], "url": folder["web_url"],
        "approval_id": state["approval_status"]["final_delivery_package"]}) for item in stories]) if stories else []
    linked = sum(isinstance(link, dict) and link.get("status") in ("linked", "already linked") for link in links)
    return folder, audit("sharepoint", f"{uploaded} documents filed in {folder['path']}; linked on {linked} User Stories")


def publish(state: DeliveryState) -> dict:
    folder = Path("output") / state["request_id"]
    folder.mkdir(parents=True, exist_ok=True)

    files = [export_design_document(state[key], state["request_id"], state["locked_facts"], folder)
             for key in ("fdd", "tdd") if state.get(key)]
    files.append(export_test_document(state["tests"], state["matrix"], state["request_id"], folder))

    cr = state.get("cr_assessment") or {}
    if cr.get("cr_required"):  # the Change Request form for the customer to sign
        context = "Immutable facts:\n" + "\n".join(state["locked_facts"]) + "\n\nPlan:\n" + plan_summary(state["plan"])
        content = write_change_request(state["source_text"], cr["change_items"], context)
        files.append(export_change_request(content.model_dump(), {
            "change_number": state["request_id"], "requested_by": state.get("requested_by", "Customer"),
            "date_of_request": state.get("request_date") or form_date(), "phase": os.getenv("CR_PROJECT_PHASE", "Phase I"),
            "effort_hours": sum(t["estimate_hours"] for s in state["plan"]["user_stories"] for t in s["tasks"]),
        }, folder))

    package_path = folder / f"{state['request_id']}_delivery_package.json"
    package_path.write_text(json.dumps({key: state.get(key) for key in (
        "request_id", "requirements", "classification", "confirmed_answers", "assumptions", "locked_facts",
        "evidence", "cr_assessment", "plan", "owner_recommendations", "fdd", "tdd", "tests", "matrix", "issues", "approval_status",
    )}, indent=2, ensure_ascii=False), encoding="utf-8")
    files.append(package_path)

    package = {"request_id": state["request_id"], "module": state["classification"]["module"],
               "work_type": state["classification"]["work_type"], "plan": state["plan"],
               "owner_recommendations": state["owner_recommendations"]}
    result = call_mcp_tool("devops", "create_linked_work_items", package_json=json.dumps(package),
                           approval_id=state["approval_status"]["final_delivery_package"])
    created = result.get("created", []) if isinstance(result, dict) else []

    log = audit("document_tool", f"{len(files)} files exported to {folder}")
    log += audit("devops", f"{len(created)} work items created ({result.get('mode', result.get('error'))})")
    sharepoint_folder, filing_log = file_on_sharepoint(state, files, created)
    return {"exported_files": [str(p) for p in files], "work_items": created, "sharepoint_folder": sharepoint_folder,
            "status": "completed", "audit_log": log + filing_log}

builder = StateGraph(DeliveryState)

builder.add_node("intake", intake)
builder.add_node("requirements_agent", requirements_agent)
builder.add_node("gap_retrieval_agent", gap_retrieval_agent)
builder.add_node("cr_assessment_agent", cr_assessment_agent)
builder.add_node("communication_agent", communication_agent)
builder.add_node("approve_clarification", approve_clarification,
                 destinations=("communication_agent", "send_clarification", "wait_for_customer", END))
builder.add_node("send_clarification", send_clarification)
builder.add_node("wait_for_customer", wait_for_customer)
builder.add_node("refine_specification", refine_specification)
builder.add_node("spec_approval", spec_approval, destinations=("refine_specification", "planning_agent", END))
builder.add_node("planning_agent", planning_agent)
builder.add_node("functional_design", design_node("FDD"))
builder.add_node("technical_design", design_node("TDD"))
builder.add_node("test_agent", test_agent)
builder.add_node("validation_agent", validation_agent,
                 destinations=("planning_agent", "functional_design", "technical_design", "test_agent", "final_approval"))
builder.add_node("final_approval", final_approval, destinations=("planning_agent", "publish", END))
builder.add_node("publish", publish)

builder.add_edge(START, "intake")
builder.add_edge("intake", "requirements_agent")
builder.add_edge("requirements_agent", "gap_retrieval_agent")
builder.add_edge("gap_retrieval_agent", "cr_assessment_agent")
builder.add_conditional_edges("cr_assessment_agent", route_after_gaps, ["communication_agent", "spec_approval"])
builder.add_edge("communication_agent", "approve_clarification")
builder.add_edge("send_clarification", "wait_for_customer")
builder.add_edge("wait_for_customer", "refine_specification")
builder.add_edge("refine_specification", "spec_approval")
builder.add_conditional_edges("planning_agent", route_after_planning,
                              ["functional_design", "technical_design", "validation_agent"])
builder.add_conditional_edges("functional_design", route_after_design, ["test_agent", "validation_agent"])
builder.add_conditional_edges("technical_design", route_after_design, ["test_agent", "validation_agent"])
builder.add_edge("test_agent", "validation_agent")
builder.add_edge("publish", END)

# run.py compiles the graph with a checkpointer: builder.compile(checkpointer=...)


# ====================================================================================================
# 10. Support emails to work items (notebook section 10)
# ====================================================================================================

import math

MAX_TASK_HOURS = 8
HOURS_PER_DAY = float(os.getenv("ADO_HOURS_PER_DAY", "6"))


class SupportTask(BaseModel):
    title: str = Field(description="Starts with FUNC: or TECH:, short and specific")
    description: str = Field(description="What to do, in one or two sentences")
    kind: Literal["Functional", "Technical"]
    activity: Literal["Requirements", "Design", "Development", "Testing", "Documentation", "Deployment"]
    estimate_hours: float = Field(description=f"Realistic estimate between 0.5 and {MAX_TASK_HOURS} hours")


class SupportTicketPlan(BaseModel):
    title: str = Field(description="The task only, without company or sender, e.g. 'Fix sales order confirmation error'")
    description: str = Field(description="The problem as bullet points, only facts from the email; mention attached files")
    module: Literal["Finance", "Supply Chain", "Finance and Supply Chain", "Out of scope"]
    priority: Literal["Low", "Medium", "High"]
    category: Literal["Bug", "Support", "Feature Request"]
    acceptance_criteria: list[str]
    steps_to_reproduce: list[str] = Field(description="Steps given in the email; empty when there are none")
    tasks: list[SupportTask]


class SupportReply(BaseModel):
    greeting_name: str = Field(description="The customer's first name")
    case_description: str = Field(description="1-2 polite sentences summarizing the issue; acknowledge attached files")
    next_action_plan: str = Field(description="1-2 sentences on how the team will investigate and resolve it")


SUPPORT_PLANNING_SYSTEM = f"""
You are the Support Planning Agent of a Dynamics 365 F&O support team.
Turn a customer support email into one Azure DevOps ticket and split the work into small tasks.

Rules:
- Use only facts from the email. Write "Not specified" for missing information and never invent values.
- The title is the task itself, without the company or the sender's name.
- Split the work into small, concrete tasks in the order they are done: investigate and reproduce, fix or
  configure, test, then document and inform the customer. Each task has one Activity and takes at most
  {MAX_TASK_HOURS} hours; split bigger work into several tasks.
- Activity: Requirements = clarifying the need with the customer, Design = solution or configuration design,
  Development = investigating, reproducing and fixing (code, configuration, data), Testing = verifying the
  fix, Documentation = documenting and informing the customer, Deployment = moving the fix to production.
- Functional tasks (analysis, configuration, testing with users) start with "FUNC:"; technical tasks
  (X++, integrations, data fixes, deployment) start with "TECH:".
- Priority is High only when the customer's business is blocked.
"""

SUPPORT_REPLY_SYSTEM = """
You write the acknowledgement email for a new Dynamics 365 F&O support ticket.
- Short, polite business language without technical jargon.
- Do not mention internal names, estimates or task IDs.
- Never promise a fix date or a solution: the next contact date is given separately.
"""


def plan_support_ticket(email: dict, feedback: str = "") -> dict:
    attachments = ", ".join(a["name"] for a in email["attachments"]) or "None"
    user_prompt = wrap(f"From: {email['from_name']} <{email['from_address']}>\nSubject: {email['subject']}\n"
                       f"Attachments: {attachments}\n\n{email['body']}")
    if feedback:
        user_prompt += f"\n\nFix this problem in your plan: {feedback}"
    return ask_structured(SupportTicketPlan, SUPPORT_PLANNING_SYSTEM, user_prompt).model_dump()


def draft_support_reply(email: dict, ticket: dict, tasks: list, feedback: str = "",
                        cr_items: list | None = None) -> SupportReply:
    user_prompt = (
        f"Customer email from {email['from_name']}:\n{wrap(email['body'])}\n\n"
        f"Attachments received: {', '.join(a['name'] for a in email['attachments']) or 'none'}\n\n"
        f"Ticket: {ticket['title']}\n{ticket['description']}\n\n"
        "Planned work:\n" + "\n".join(f"- {t['title']}" for t in tasks)
        + (f"\n\nWe are sending the customer a Change Request (written by us, in its own paragraph) for: "
           f"{'; '.join(cr_items)}. Do not describe it as something the customer sent." if cr_items else "")
    )
    if feedback:
        user_prompt += f"\n\nReviewer feedback on the previous draft: {feedback}"
    return ask_structured(SupportReply, SUPPORT_REPLY_SYSTEM, user_prompt)


def ticket_title(customer: str, title: str) -> str:
    """'Customer - short task', without repeating the customer name the model may have added."""
    title = title.strip()
    while title.lower().startswith(customer.lower() + " - "):
        title = title[len(customer) + 3:].strip()
    return f"{customer} - {title or 'Email-based support request'}"


def add_business_days(day: date, days: int) -> date:
    while day.weekday() >= 5:  # start on a working day
        day += timedelta(days=1)
    for _ in range(days):
        day += timedelta(days=1)
        while day.weekday() >= 5:
            day += timedelta(days=1)
    return day


def schedule_tasks(tasks: list, start: date) -> list:
    """Tasks are done one after the other: each due date is the business day on which its work is finished."""
    done, scheduled = 0.0, []
    for number, task in enumerate(tasks, start=1):
        done += task["estimate_hours"]
        due = add_business_days(start, max(math.ceil(done / HOURS_PER_DAY) - 1, 0))
        scheduled.append({**task, "task_id": f"T-{number:02d}", "due_date": due.isoformat()})
    return scheduled


def render_support_reply(reply: SupportReply, ticket_id, next_contact: str, cr: dict | None = None) -> str:
    return (
        f"Dear {reply.greeting_name},\n\n"
        f"This is to inform you that a support ticket **#{ticket_id}** has been created and dispatched. "
        "Please find the details below:\n\n"
        f"**Case Description:**\n{reply.case_description}\n\n"
        f"**Next Action Plan:**\n{reply.next_action_plan}\n\n"
        + cr_reply_paragraph(cr) +
        "**Scope Agreement:**\nThis case will be considered resolved and archived once the issue is fully "
        "resolved and required access is restored (if applicable).\n\n"
        "Should any new or additional issues arise during the resolution of this incident, a separate support "
        "ticket must be created.\n\n"
        f"**Next Contact Date:**\nOn or before {next_contact}.\n\n"
        f"Thank you for your attention and cooperation.\nBest regards,\n{APPROVER}"
    )

class SupportState(TypedDict, total=False):
    message_id: str
    email: dict
    feature: dict
    customer: str
    domain: str
    ticket: dict
    cr: dict
    tasks: list
    change_request: dict
    attachment_paths: list
    sharepoint_folder: dict
    ticket_approval: str
    created: dict
    reply: dict
    reply_feedback: str
    reply_draft: dict
    status: str
    security_findings: Annotated[list, operator.add]  # prompt-injection findings, shown at every gate
    audit_log: Annotated[list, operator.add]


def read_support_email(state: SupportState) -> dict:
    email = call_mcp_tool("email", "get_email", message_id=state["message_id"])
    if "error" in email:
        raise RuntimeError(email["error"])
    email["body"], findings = screen(email["body"])  # input check: hidden text removed, instruction-like text flagged
    email["subject"], subject_findings = screen(email["subject"])
    findings += subject_findings
    log = audit("supervisor", f"support email '{email['subject']}' from {email['from_address']} "
                              f"({len(email['attachments'])} attachments)")
    if findings:
        log += audit("guard", f"possible prompt injection in the email: {', '.join(f['rule'] for f in findings)}")
    return {"email": email, "security_findings": [{**f, "source": "support email"} for f in findings], "audit_log": log}


def find_feature(state: SupportState) -> dict:
    suggestion = call_mcp_tool("devops", "suggest_features_for_sender", sender_email=state["email"]["from_address"])
    features = suggestion["features"][:15]
    if not features:
        raise ValueError("No support Features found in Azure DevOps.")
    default = features.index(suggestion["default"]) if suggestion["default"] in features else None
    choice = interrupt({"type": "choice", "title": f"Support Feature for '{suggestion['domain']}'", "default": default,
                        "options": [f"#{f['id']} {f['title']}  [{f['project']}]" for f in features]})
    feature = features[choice["index"]]

    if suggestion["from_memory"]:
        how = "remembered for this domain"
    elif suggestion["match"]["project"] and feature["project"] == suggestion["match"]["project"]:
        how = f"matched by {suggestion['match']['matched_by']}"
    else:
        how = "chosen by a person"
    return {"feature": feature, "customer": feature["title"].split(" - ")[0], "domain": suggestion["domain"],
            "audit_log": audit("devops", f"Feature #{feature['id']} {feature['title']} ({how})")}


def support_planning_agent(state: SupportState) -> dict:
    ticket = plan_support_ticket(state["email"])
    too_big = [t["title"] for t in ticket["tasks"] if t["estimate_hours"] > MAX_TASK_HOURS]
    if too_big or not ticket["tasks"]:
        ticket = plan_support_ticket(state["email"], feedback=(
            f"Split these tasks into tasks of at most {MAX_TASK_HOURS} hours: {too_big}" if too_big else "Add the tasks."))
    ticket["title"] = ticket_title(state["customer"], ticket["title"])
    hours = sum(t["estimate_hours"] for t in ticket["tasks"])
    return {"ticket": ticket, "audit_log": audit("support_planning", f"{ticket['title']}: {len(ticket['tasks'])} tasks, "
                                                                     f"{hours} hours, {ticket['priority']} {ticket['category']}")}


def assess_cr(state: SupportState) -> dict:
    email = state["email"]
    assessment = assess_change_request(f"Subject: {email['subject']}\n\n{email['body']}",
                                       context=f"Ticket category from the Support Planning Agent: {state['ticket']['category']}")
    verdict = "CR required" if assessment.cr_required else "no CR needed"
    return {"cr": assessment.model_dump(), "audit_log": audit("cr_assessment", f"{assessment.request_type}: {verdict}")}


def schedule_and_owners(state: SupportState) -> dict:
    tasks = schedule_tasks(state["ticket"]["tasks"], date.today())
    workload = call_mcp_tool("devops", "get_team_workload", project=state["feature"]["project"] if ADO_LIVE else "")
    workload = workload if isinstance(workload, dict) and "error" not in workload else {}
    owners = recommend_owners({"user_stories": [{"tasks": tasks}]}, state["ticket"]["module"], TEAM, workload,
                              tasks[-1]["due_date"])
    for task in tasks:
        task["owner"], task["owner_reason"] = owners[task["task_id"]]["owner"], owners[task["task_id"]]["reason"]
    return {"tasks": tasks, "audit_log": audit("supervisor", f"due dates {tasks[0]['due_date']} to {tasks[-1]['due_date']}, "
                                                             "owners recommended")}


def approve_ticket(state: SupportState) -> Command:
    feature, ticket = state["feature"], state["ticket"]
    artefact = (
        f"Feature: #{feature['id']} {feature['title']} ({feature['project']})\n"
        f"User Story: {ticket['title']}  [{ticket['priority']}, {ticket['category']}, {ticket['module']}]\n"
        + "\n".join(f"  {t['task_id']} {t['title']} | {t['activity']} | {t['estimate_hours']}h | due {t['due_date']} | "
                    f"{t['owner']}" for t in state["tasks"])
        + f"\nAttachments: {', '.join(a['name'] for a in state['email']['attachments']) or 'none'}"
        + f"\nChange Request: {'REQUIRED (' + state['cr']['request_type'] + ')' if state['cr']['cr_required'] else 'not needed'}"
        + "".join(f"\n   - {reason}" for reason in state["cr"]["reasons"])
    )
    review = interrupt(guarded(state, {"type": "approval", "title": "Support ticket to create in Azure DevOps",
                                       "artefact": artefact, "options": ["approve", "reject"], "live": ADO_LIVE}))
    decision = DECISIONS[review["decision"]]
    record = approval_ledger.record_decision("support_ticket", state["message_id"], APPROVER, decision, review["comments"])
    log = audit("human", f"support ticket {decision} ({record['approval_id']})")
    if decision != "approved":
        return Command(goto=END, update={"status": "stopped: support ticket rejected", "audit_log": log})
    return Command(goto="create_ticket", update={"ticket_approval": record["approval_id"], "audit_log": log})


def create_ticket(state: SupportState) -> dict:
    files, log = [], []
    if state["email"]["attachments"]:
        files = call_mcp_tool("email", "save_attachments", message_id=state["message_id"])
        if isinstance(files, dict):  # tool error: create the ticket without the files
            log += audit("email", f"attachments not downloaded: {files['error']}")
            files = []
    package = {"message_id": state["message_id"], "domain": state["domain"], "customer": state["customer"],
               "email_subject": state["email"]["subject"], "feature": state["feature"], "ticket": state["ticket"],
               "tasks": state["tasks"], "attachment_paths": [f["path"] for f in files],
               "cr_required": state["cr"]["cr_required"]}
    created = call_mcp_tool("devops", "create_support_work_items", package_json=json.dumps(package),
                            approval_id=state["ticket_approval"])
    if "error" in created:
        raise RuntimeError(created["error"])
    attached = sum(a["status"] == "attached" for a in created["attachments"])
    log += audit("devops", f"User Story #{created['story']['id']} + {len(created['tasks'])} tasks, "
                           f"{attached}/{len(created['attachments'])} attachments ({created['mode']})")
    return {"created": created, "attachment_paths": package["attachment_paths"], "audit_log": log}


def route_after_ticket(state: SupportState) -> str:
    return "change_request_document" if state["cr"]["cr_required"] else "file_to_sharepoint"


def change_request_document(state: SupportState) -> dict:
    """The Change Request form for the changed parts, attached to the new User Story."""
    email, story_id = state["email"], state["created"]["story"]["id"]
    content = write_change_request(
        f"Subject: {email['subject']}\n\n{email['body']}", state["cr"]["change_items"],
        f"Ticket: {state['ticket']['title']}\n{state['ticket']['description']}\n\nPlanned tasks:\n"
        + "\n".join(f"- {t['title']} ({t['estimate_hours']}h)" for t in state["tasks"]),
    )
    path = export_change_request(content.model_dump(), {
        "change_number": story_id, "requested_by": f"{email['from_name']} ({state['customer']})",
        "date_of_request": form_date(email.get("received", "")), "phase": os.getenv("CR_SUPPORT_PHASE", "Support"),
        "effort_hours": sum(t["estimate_hours"] for t in state["tasks"]),
    }, Path("output") / "support" / str(story_id))
    attached = call_mcp_tool("devops", "attach_documents", work_item_id=story_id, file_paths=[str(path.resolve())],
                             message_id=state["message_id"], approval_id=state["ticket_approval"])
    status = attached[0]["status"] if isinstance(attached, list) and attached else attached
    return {"change_request": {"number": f"CR-{story_id}", "request_type": state["cr"]["request_type"],
                               "path": str(path.resolve()), "name": content.change_name},
            "audit_log": audit("change_request", f"{path.name} written; on User Story #{story_id}: {status}")}


def file_to_sharepoint(state: SupportState) -> dict:
    """File the customer's attachments and the CR as <root>/<Customer>/<Story ID - Description>/<type>/, link the folder."""
    story_id = state["created"]["story"]["id"]
    folder_args = {"customer": state["customer"], "item_id": str(story_id), "description": state["ticket"]["title"]}
    calls = []
    if state.get("attachment_paths"):
        calls.append(("file_documents", {**folder_args, "doc_type": "Customer Files", "file_paths": state["attachment_paths"]}))
    if state.get("change_request"):
        calls.append(("file_documents", {**folder_args, "doc_type": "Change Request",
                                         "file_paths": [state["change_request"]["path"]]}))
    *filings, folder = call_mcp_tools("sharepoint", calls + [("ticket_folder_url", folder_args)])
    if "error" in folder:
        raise RuntimeError(folder["error"])
    link = call_mcp_tool("devops", "add_document_link", work_item_id=story_id, url=folder["web_url"],
                         approval_id=state["ticket_approval"], message_id=state["message_id"])
    uploaded = sum(f["status"] == "uploaded" for r in filings if isinstance(r, dict) for f in r.get("files", []))
    return {"sharepoint_folder": folder, "audit_log": audit(
        "sharepoint", f"{uploaded} documents filed in {folder['path']}; folder linked on #{story_id}: {link.get('status', link)}")}


def support_reply_agent(state: SupportState) -> dict:
    cr = state.get("change_request")
    reply = draft_support_reply(state["email"], state["ticket"], state["tasks"], state.get("reply_feedback", ""),
                                state["cr"]["change_items"] if cr else None)
    ticket_id = state["created"]["story"]["id"]
    next_contact = add_business_days(date.today(), 1).strftime("%B %d, %Y")
    return {"reply": {"subject": f"{state['ticket']['title']} | {ticket_id}",
                      "body": render_support_reply(reply, ticket_id, next_contact, cr)},
            "audit_log": audit("support_reply", f"acknowledgement drafted for ticket #{ticket_id}")}


def approve_reply(state: SupportState) -> Command:
    reply = state["reply"]
    review = interrupt(guarded(state, {
        "type": "approval", "title": "Acknowledgement email to the customer",
        "artefact": f"Subject: {reply['subject']}\n\n{reply['body']}\n\nAttachments: "
                    f"{Path(state['change_request']['path']).name if state.get('change_request') else 'none'}"},
        outgoing=reply["body"]))
    decision = DECISIONS[review["decision"]]
    record = approval_ledger.record_decision("customer_email", state["message_id"], APPROVER, decision, review["comments"])
    log = audit("human", f"acknowledgement email {decision} ({record['approval_id']})")
    if decision == "revise":
        return Command(goto="support_reply_agent", update={"reply_feedback": review["comments"], "audit_log": log})
    if decision == "rejected":
        return Command(goto=END, update={"status": "stopped: acknowledgement rejected (ticket created)", "audit_log": log})

    draft = call_mcp_tool("email", "save_reply_draft", message_id=state["message_id"], subject=reply["subject"],
                          body=reply["body"], approval_id=record["approval_id"],
                          attachment_paths=[state["change_request"]["path"]] if state.get("change_request") else [])
    if "error" in draft:
        raise RuntimeError(draft["error"])
    log += audit("email", f"reply draft saved ({draft['mode']}, {draft['status']})")
    return Command(goto="send_reply", update={"reply_draft": draft, "audit_log": log})


def send_reply(state: SupportState) -> dict:
    return {**send_gate(state["reply_draft"], "acknowledgement email", state), "status": "completed"}

support_builder = StateGraph(SupportState)
support_builder.add_node("read_support_email", read_support_email)
support_builder.add_node("find_feature", find_feature)
support_builder.add_node("support_planning_agent", support_planning_agent)
support_builder.add_node("assess_cr", assess_cr)
support_builder.add_node("schedule_and_owners", schedule_and_owners)
support_builder.add_node("approve_ticket", approve_ticket, destinations=("create_ticket", END))
support_builder.add_node("create_ticket", create_ticket)
support_builder.add_node("change_request_document", change_request_document)
support_builder.add_node("file_to_sharepoint", file_to_sharepoint)
support_builder.add_node("support_reply_agent", support_reply_agent)
support_builder.add_node("approve_reply", approve_reply, destinations=("support_reply_agent", "send_reply", END))
support_builder.add_node("send_reply", send_reply)

support_builder.add_edge(START, "read_support_email")
support_builder.add_edge("read_support_email", "find_feature")
support_builder.add_edge("find_feature", "support_planning_agent")
support_builder.add_edge("support_planning_agent", "assess_cr")
support_builder.add_edge("assess_cr", "schedule_and_owners")
support_builder.add_edge("schedule_and_owners", "approve_ticket")
support_builder.add_conditional_edges("create_ticket", route_after_ticket,
                                      ["change_request_document", "file_to_sharepoint"])
support_builder.add_edge("change_request_document", "file_to_sharepoint")
support_builder.add_edge("file_to_sharepoint", "support_reply_agent")
support_builder.add_edge("support_reply_agent", "approve_reply")
support_builder.add_edge("send_reply", END)

# run.py compiles the graph with a checkpointer: support_builder.compile(checkpointer=...)
