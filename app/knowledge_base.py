"""Knowledge base (notebook section 2): the approved knowledge, chunked, embedded with nomic-embed-text (Ollama)
and saved to fo_knowledge_store.json for the knowledge MCP server.

Drop your own .md / .txt / .docx files into knowledge/ (in the data folder) and run `python run.py setup` again.
"""
import os
from pathlib import Path

import docx
from langchain_core.documents import Document
from langchain_core.vectorstores import InMemoryVectorStore
from langchain_ollama import OllamaEmbeddings
from langchain_text_splitters import RecursiveCharacterTextSplitter

KNOWLEDGE_DIR = Path("knowledge")
STORE_FILE = Path("fo_knowledge_store.json")
# Summarised notes used as the organization's approved knowledge.
# source_type: microsoft_guidance | company_standard | approved_design | email_template
KNOWLEDGE_DOCS = [
    # ---------------- Microsoft guidance (summarised) ----------------
    {"doc_id": "MS-EXT-01", "title": "Extensibility model and Chain of Command", "source_type": "microsoft_guidance", "module": "General", "approved": True,
     "text": "Dynamics 365 Finance and Operations is customized through extensions. Overlayering of Microsoft models is not supported, so standard objects are changed only through extension objects placed in a separate model. Tables, forms, data entities, enums and menus are extended with extension elements, for example a table extension that adds new fields. Business logic is extended with Chain of Command (CoC): a final extension class decorated with ExtensionOf wraps a public or protected method and must call next to run the standard logic. Event handlers (pre/post events and data events such as onValidatedWrite) are an alternative for reacting to standard events. Extensions keep the solution upgradable through Microsoft's continuous updates."},
    {"doc_id": "MS-SEC-01", "title": "Role-based security and segregation of duties", "source_type": "microsoft_guidance", "module": "General", "approved": True,
     "text": "Security in F&O is role based. Users are assigned security roles, roles contain duties, duties contain privileges, and privileges grant access levels to entry points such as menu items, and to tables and data entities. New menu items, forms or actions require new privileges, which are grouped into duties and added to roles through extensions rather than by changing standard roles. Segregation of duties (SoD) rules in System administration detect or prevent a user from holding conflicting duties, for example entering and approving vendor invoices. Workflow approvals should be assigned to roles or positions instead of named users so they keep working when people change."},
    {"doc_id": "MS-WF-01", "title": "Workflow and approvals", "source_type": "microsoft_guidance", "module": "Finance", "approved": True,
     "text": "F&O provides configurable approval workflows for documents such as vendor invoices, purchase requisitions, purchase orders and general journals. A workflow is configured per legal entity in the relevant module, for example Accounts payable > Setup > Accounts payable workflows. Approval steps can have conditions, such as the invoice total being greater than an amount, so that only documents above a threshold require approval. Assignments can target a user, a role, a position or a hierarchy. Each step supports time limits and escalation paths, for example escalate to another user after a number of days, and optional delegation. The workflow history records who submitted, approved or rejected a document and when, which supports audit requirements. When the vendor invoice workflow is active, invoices cannot be posted until the workflow is approved."},
    {"doc_id": "MS-AP-01", "title": "Vendor invoices and holds", "source_type": "microsoft_guidance", "module": "Finance", "approved": True,
     "text": "Vendor invoices are registered in the vendor invoice journal or the pending vendor invoices page, or created from purchase orders. Invoice matching validates invoices against purchase orders and product receipts (two-way or three-way matching) within configured price and quantity tolerances. A vendor can be placed on hold with the vendor's On hold field: Invoice blocks invoice entry and posting, Payment blocks payment generation while invoices can still be registered, and All blocks all transactions. Vendor bank accounts are maintained on the vendor record, and one of them can be selected as the vendor's default bank account. Changes to vendor master data can be controlled with the vendor approval (proposed changes) feature."},
    {"doc_id": "MS-GL-01", "title": "General ledger foundations", "source_type": "microsoft_guidance", "module": "Finance", "approved": True,
     "text": "The general ledger uses a chart of accounts shared across legal entities, main accounts and financial dimensions to build ledger accounts. Posting profiles map subledger transactions to main accounts. Ledger calendars and fiscal periods control which periods are open for posting. Number sequences provide document numbers per legal entity and should be configured rather than hard-coded."},
    {"doc_id": "MS-DATA-01", "title": "Data management and integration", "source_type": "microsoft_guidance", "module": "General", "approved": True,
     "text": "The data management framework uses data entities for import and export. Data projects can run once or as recurring integrations, and data packages can move configuration between environments. OData endpoints expose public data entities for lightweight real-time integration; large volumes should use the package API or recurring integrations. A custom data entity needs a staging table and is created with the data entity wizard in Visual Studio."},
    {"doc_id": "MS-BATCH-01", "title": "Batch processing with SysOperation", "source_type": "microsoft_guidance", "module": "General", "approved": True,
     "text": "Long-running or scheduled logic should run in the batch framework instead of on a form. The SysOperation framework separates a data contract (parameters), a controller (how the operation runs) and a service class (the business logic). Operations can be scheduled as recurring batch jobs, assigned to batch groups and monitored in System administration > Batch jobs. Batch logic should be idempotent so that a rerun does not create duplicate changes, and it should log the records it changes."},
    {"doc_id": "MS-PROC-01", "title": "Procurement and purchase order approval", "source_type": "microsoft_guidance", "module": "Supply Chain", "approved": True,
     "text": "Purchase requisitions let employees request goods and services; approved requisitions are released to purchase orders. Procurement categories classify purchases and drive policies. Change management on purchase orders activates the purchase order approval workflow, so a purchase order must be approved before it is confirmed and later changes require re-approval. Purchase agreements define committed quantities or amounts with vendors."},
    {"doc_id": "MS-INV-01", "title": "Inventory dimensions and batch control", "source_type": "microsoft_guidance", "module": "Supply Chain", "approved": True,
     "text": "Items are controlled by item model groups (costing method and inventory policies) and dimension groups. Storage dimensions include site, warehouse and location; tracking dimensions include batch and serial numbers. Batch-controlled items can store expiration dates, and inventory statuses or inventory blocking can prevent the use of specific on-hand quantities. Inventory closing settles issue and receipt transactions for the selected costing method."},

    # ---------------- Company standards ----------------
    {"doc_id": "STD-DEV-01", "title": "X++ development standard", "source_type": "company_standard", "module": "General", "approved": True,
     "text": "All custom objects use the prefix CUS and are placed in the model CUSExtensions. Customizations must use extensions and Chain of Command; overlayering and copying standard code are not allowed. Configurable values such as thresholds, days or amounts are stored in a parameter table exposed on a parameters form, never hard-coded. Every new menu item needs a privilege and a duty, delivered as extensions of existing roles. Business logic must have SysTest unit tests, and each change goes through pull-request code review in Azure DevOps."},
    {"doc_id": "STD-DOC-01", "title": "Design document standard", "source_type": "company_standard", "module": "General", "approved": True,
     "text": "A Functional Design Document (FDD) is required when a business process or configuration changes. It contains business context, as-is and to-be process, configuration and parameters, roles and security, exceptions, reports and data, assumptions and open points. A Technical Design Document (TDD) is required when code, integrations or reports are created. It contains solution overview, object list (extensions, classes, tables, security objects), logic, data model changes, error handling and logging, performance and batch considerations, and deployment steps. Mixed changes need both documents. Every section references the requirement IDs it covers."},
    {"doc_id": "STD-ADO-01", "title": "Azure DevOps work item standard", "source_type": "company_standard", "module": "General", "approved": True,
     "text": "Each approved requirement is delivered through at least one User Story. Story titles use the format '[Module] short goal', and the description uses 'As a <role>, I want <goal> so that <benefit>'. Acceptance criteria are written as Given/When/Then and reference requirement IDs. Implementation work is split into child Tasks: functional tasks are prefixed 'FUNC:' (configuration, FDD, training) and technical tasks are prefixed 'TECH:' (development, TDD, unit tests). Tasks carry an estimate in hours in Remaining Work. Stories are tagged with the module, work type and request ID. Assignees are proposed and then confirmed by a person; tools must not assign work automatically."},
    {"doc_id": "STD-TEST-01", "title": "Test case standard", "source_type": "company_standard", "module": "General", "approved": True,
     "text": "Each requirement is covered by at least one test case, and test cases reference requirement IDs. Test types: Functional (business scenario in the UI), Technical (unit or component behaviour, such as a SysTest), Integration (interaction with workflow, batch, data entities or external systems) and Regression (standard processes that must still work). Threshold rules need boundary tests at, just below and just above the limit. Security requirements need a negative test proving that an unauthorized user is blocked."},
    {"doc_id": "STD-COMM-01", "title": "Customer communication standard", "source_type": "company_standard", "module": "General", "approved": True,
     "text": "Clarification emails are short, polite and written in business language without technical jargon. Questions are numbered, one topic per question, with at most eight questions per email. Each email restates the request in one sentence, asks for a reply date, and never promises a delivery date or solution before the requirement is approved. Emails are saved as drafts and sent only after a person reviews them."},
    {"doc_id": "STD-CR-01", "title": "Change request standard", "source_type": "company_standard", "module": "General", "approved": True,
     "text": "A Change Request (CR) is required before any work that changes the agreed solution: new or changed functionality, customizations (X++ extensions), new or changed reports, integrations, data entities, workflows or fields, changes to the agreed scope, legal entities or business processes, and any configuration change outside the support contract. No CR is needed for support: defects in functionality that was delivered and accepted, errors caused by data or setup that are restored to the agreed design, how-to questions, access or security issues within the agreed roles, and standard Microsoft updates. When a request mixes both, the support part is handled as a ticket and the changed part as a CR. A CR uses the ERP Change/Additional Request Form: project name, change number, requested by, date of request, required phase, business priority, change name and technical complexity; the change description (business requirement and proposed solution) and the recommendation/best practice (customization required, affected areas, estimated effort) are written by the implementation partner; the business justification and the impact of not implementing the change are left for the customer's business to fill in; sign-off by the customer's functional lead, head of department and IT manager and by the partner's project manager. Work on a change starts only after the CR is signed."},

    # ---------------- Previously approved designs ----------------
    {"doc_id": "DES-2025-014", "title": "Approved design: purchase invoice approval by amount", "source_type": "approved_design", "module": "Finance", "approved": True,
     "text": "The vendor invoice workflow was activated per legal entity with the condition 'Invoice total (accounting currency) > threshold'. Invoices at or below the threshold were auto-approved with an automatic action so AP clerks could post them directly. The approval step was assigned to the Accounts payable manager role, with escalation to the finance director after 3 days. Segregation of duties was enforced with an SoD rule between the duties 'Maintain vendor invoices' and 'Approve vendor invoices', and the workflow configuration prevented the submitter from approving. The threshold was stored in a CUS parameter on the Accounts payable parameters form. Audit used the standard workflow history."},
    {"doc_id": "DES-2025-022", "title": "Approved design: capital purchase order approval", "source_type": "approved_design", "module": "Supply Chain", "approved": True,
     "text": "Purchase orders for procurement categories flagged as capital expenditure required approval by the plant controller. Change management was enabled for the affected legal entities, and the purchase order workflow used a condition on the procurement category. A CUS field on the procurement category table flagged capital categories. Regression tests covered non-capital purchase orders and purchase order changes after confirmation."},
    {"doc_id": "DES-2025-031", "title": "Approved design: automatic hold for vendors with incomplete master data", "source_type": "approved_design", "module": "Finance", "approved": True,
     "text": "A recurring SysOperation batch job (CUSVendMasterDataCheck) evaluated vendors in selected legal entities, set On hold = Payment when mandatory data was missing, and released the hold when the data was completed unless the hold had been set manually. A Chain of Command extension on the vendor table validated new vendors on insert and update. Each change was logged in a CUS log table with vendor, legal entity, old value, new value, reason and timestamp. The batch was idempotent and processed vendors per legal entity."},
    {"doc_id": "DES-DRAFT-007", "title": "Draft design (not approved): auto-post small invoices", "source_type": "approved_design", "module": "Finance", "approved": False,
     "text": "Draft idea to post vendor invoices below 500 USD automatically without any review. Rejected by the design authority and must not be reused."},

    # ---------------- Email templates ----------------
    {"doc_id": "TPL-CLARIFY-01", "title": "Clarification request template", "source_type": "email_template", "module": "General", "approved": True,
     "text": "Subject: Clarification needed - {request title}\n\nDear {customer name},\n\nThank you for your request regarding {one-sentence summary}. To make sure we deliver exactly what you need, could you please help us with the following questions:\n\n{numbered questions}\n\nCould you please reply by {reply date}? Once we have your answers, we will confirm the final scope before starting the work.\n\nKind regards,\n{consultant name}\nDynamics 365 F&O Delivery Team"},
    {"doc_id": "TPL-CONFIRM-01", "title": "Scope confirmation template", "source_type": "email_template", "module": "General", "approved": True,
     "text": "Subject: Scope confirmation - {request title}\n\nDear {customer name},\n\nThank you for your answers. Please find below the agreed requirements:\n\n{requirements}\n\nAssumptions:\n{assumptions}\n\nPlease confirm that this scope is correct so that we can start the design.\n\nKind regards,\n{consultant name}\nDynamics 365 F&O Delivery Team"},
]


def write_knowledge_files(folder: Path = KNOWLEDGE_DIR) -> None:
    """Each document becomes a Markdown file with a small metadata header (your own files are kept)."""
    folder.mkdir(exist_ok=True)
    for doc in KNOWLEDGE_DOCS:
        header = "\n".join(f"{key}: {doc[key]}" for key in ("doc_id", "title", "source_type", "module", "approved"))
        (folder / f"{doc['doc_id']}.md").write_text(f"---\n{header}\n---\n{doc['text']}\n", encoding="utf-8")


def load_knowledge_folder(folder: Path) -> list[Document]:
    """Load .md/.txt/.docx files. Metadata comes from the header; files without one count as approved company standards."""
    documents = []
    for path in sorted(folder.iterdir()):
        if path.suffix.lower() in (".md", ".txt"):
            raw = path.read_text(encoding="utf-8")
        elif path.suffix.lower() == ".docx":
            raw = "\n".join(p.text for p in docx.Document(path).paragraphs if p.text.strip())
        else:
            continue

        metadata = {"doc_id": path.stem, "title": path.stem, "source_type": "company_standard",
                    "module": "General", "approved": True}
        if raw.startswith("---"):
            _, header, raw = raw.split("---", 2)
            for line in header.strip().splitlines():
                key, value = line.split(":", 1)
                metadata[key.strip()] = value.strip()
            metadata["approved"] = str(metadata["approved"]).lower() == "true"

        documents.append(Document(page_content=raw.strip(), metadata=metadata))
    return documents


def build_knowledge_store() -> int:
    """Write the knowledge files, chunk and embed them, and save the store. Returns the number of chunks."""
    write_knowledge_files()
    splitter = RecursiveCharacterTextSplitter(chunk_size=700, chunk_overlap=100)
    chunks = splitter.split_documents(load_knowledge_folder(KNOWLEDGE_DIR))
    for doc_id in {c.metadata["doc_id"] for c in chunks}:
        for index, chunk in enumerate(c for c in chunks if c.metadata["doc_id"] == doc_id):
            chunk.metadata["chunk"] = index

    vector_store = InMemoryVectorStore(OllamaEmbeddings(model=os.getenv("EMBED_MODEL", "nomic-embed-text")))
    vector_store.add_documents(chunks)
    vector_store.dump(str(STORE_FILE))  # the knowledge MCP server loads it without re-embedding everything
    return len(chunks)